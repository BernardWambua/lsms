"""
The Training Needs Assessment (TNA) report, built from the gap engine.

The report follows the shape Kenyan public institutions use for TNA: an
organisational, occupational (job role) and individual analysis, a ranked list
of training needs, and the training projections (the officers to train per
skill) that HRM&D submits to the Human Resource Management Advisory Committee.
Every figure comes from GapEngine; nothing here re-derives a gap.

``build_report`` returns plain data. ``render_pdf`` and ``render_xlsx`` turn
that data into the two downloadable formats.
"""

import io
import re
from collections import defaultdict

from django.utils import timezone

from .gaps import HIGH_PRIORITY, MEDIUM_PRIORITY, priority_label
from .models import PROFICIENCY_LEVELS

LEVEL_NAMES = dict(PROFICIENCY_LEVELS)
PRIORITY_NAMES = {'high': 'High', 'medium': 'Medium', 'low': 'Low', 'none': 'Met', 'not_assessed': 'Not assessed'}
# Below this share of required skills rated, the report warns that the assessment is incomplete.
MIN_COVERAGE = 80

PERIOD_PATTERN = re.compile(r'^\d{4}/\d{2}$')


def financial_year(today=None):
    """The Kenyan government financial year (July to June) that ``today`` falls in, e.g. '2026/27'."""
    today = today or timezone.localdate()
    start = today.year if today.month >= 7 else today.year - 1
    return f'{start}/{(start + 1) % 100:02d}'


def _name(user_ref):
    return user_ref['name'] or user_ref['username']


def skill_need(skill):
    """A training needs row: one skill across the group, banded by its average priority per officer assessed."""
    mean_priority = skill['priority_total'] / skill['staff_assessed'] if skill['staff_assessed'] else None
    return {
        'skill': skill['skill'],
        'staff_requiring': skill['staff_requiring'],
        'staff_assessed': skill['staff_assessed'],
        'officers_to_train': skill['staff_with_gap'],
        'high_priority': skill['high_priority'],
        'average_gap': skill['average_gap'],
        'priority_total': skill['priority_total'],
        'priority': priority_label(None if mean_priority is None else round(mean_priority, 2)),
    }


def build_report(engine, unit=None, period=None):
    """The TNA report for the whole institution, or for ``unit`` and everything below it."""
    if unit is None:
        group = engine.institution()
        people = engine.people
        scope = {'type': 'institution', 'org_unit': None}
        roots = [engine.units[u['org_unit']['id']] for u in group['units']]
        without_role = group['staff_without_role']
    else:
        group = engine.department(unit)
        people = engine.people_in(engine.subtree(unit))
        scope = {'type': 'department', 'org_unit': group['org_unit']}
        roots = [unit]
        without_role = 0

    people_by_id = {user.pk: user for user in people}
    individuals, projections = [], []
    for summary in group['people']:
        person = engine.person(people_by_id[summary['user']['id']])
        gaps = [row for row in person['skills'] if row['gap']]
        not_assessed = [row['skill']['name'] for row in person['skills'] if row['gap'] is None]
        individuals.append({**summary, 'gaps': gaps, 'not_assessed': not_assessed})
        for row in gaps:
            projections.append({
                'skill': row['skill'],
                'officer': summary['user'],
                'job_role': summary['job_role'],
                'org_unit': summary['org_unit'],
                'required_level': row['required_level'],
                'current_level': row['current_level'],
                'rating_source': row['rating_source'],
                'gap': row['gap'],
                'priority_score': row['priority_score'],
                'priority': row['priority'],
            })
    projections.sort(key=lambda p: (-p['priority_score'], p['skill']['name'], p['officer']['username']))

    roles = defaultdict(list)
    for user in people:
        if user.job_role_id:
            roles[user.job_role_id].append(user)
    occupational = []
    for role_people in roles.values():
        role = role_people[0].job_role
        role_group = engine.summarise(role_people)
        occupational.append({
            'job_role': {'id': role.pk, 'title': role.title},
            'org_unit': {'id': role.org_unit.pk, 'name': role.org_unit.name} if role.org_unit else None,
            'staff': role_group['staff'],
            'staff_assessed': role_group['staff_assessed'],
            'staff_with_gaps': role_group['staff_with_gaps'],
            'assessment_coverage': role_group['assessment_coverage'],
            'gap_score': role_group['gap_score'],
            'priority_needs': [s['skill']['name'] for s in role_group['skills'] if s['staff_with_gap']][:5],
        })
    occupational.sort(key=lambda r: (r['gap_score'] is None, -(r['gap_score'] or 0), r['job_role']['title']))

    needs = [skill_need(skill) for skill in group['skills']]
    summary = {
        'staff': group['staff'],
        'staff_assessed': group['staff_assessed'],
        'staff_with_gaps': group['staff_with_gaps'],
        'staff_without_role': without_role,
        'assessment_coverage': group['assessment_coverage'],
        'gap_score': group['gap_score'],
        'skills_with_gaps': sum(1 for need in needs if need['officers_to_train']),
        'officers_to_train': len({p['officer']['id'] for p in projections}),
        'high_priority_gaps': sum(1 for p in projections if p['priority'] == 'high'),
    }
    return {
        'title': 'Training Needs Assessment Report',
        'institution': {'id': engine.tenant.pk, 'name': engine.tenant.name},
        'scope': scope,
        'period': period or financial_year(),
        'generated_at': timezone.now().isoformat(timespec='seconds'),
        'summary': summary,
        'methodology': methodology(),
        'organisational': organisational(engine, roots),
        'occupational': occupational,
        'training_needs': needs,
        'individuals': individuals,
        'projections': projections,
        'recommendations': recommendations(summary, needs),
    }


def organisational(engine, roots):
    """Every unit under ``roots`` in tree order, each rolling up the units below it, with its depth for indenting."""
    rows, seen = [], set()

    def walk(unit, depth):
        if unit.pk in seen:
            return
        seen.add(unit.pk)
        rows.append({**engine.unit_overview(unit), 'depth': depth})
        for child in engine.children.get(unit.pk, []):
            walk(child, depth + 1)

    for root in roots:
        walk(root, 0)
    return rows


def methodology():
    return [
        'Each job role has a skill profile setting the proficiency level the role requires (1 Awareness to 5 Expert) '
        'and how difficult, important and frequently used each skill is in that role (the DIF scores, 1 to 5).',
        'Officers rated their own proficiency through self-assessment questionnaires; supervisors then rated them. '
        'The latest supervisor rating is used where one exists, otherwise the latest self rating.',
        'Gap = required level minus current level (never below zero). Priority = gap x the mean DIF score: '
        f'{HIGH_PRIORITY} and above is high, {MEDIUM_PRIORITY} and above is medium, anything above zero is low.',
        'The gap score (0 to 100) is the DIF-weighted shortfall as a share of the DIF-weighted required levels. '
        '0 means every requirement is met.',
        'Skills nobody has rated yet are reported as not assessed and left out of all scores.',
    ]


def recommendations(summary, needs):
    items = []
    top = [need for need in needs if need['officers_to_train']][:3]
    if top:
        listed = '; '.join(f"{need['skill']['name']} ({need['officers_to_train']} officer{'s' if need['officers_to_train'] != 1 else ''})" for need in top)
        items.append(f'Prioritise training in {listed} in the training projections for this period.')
    if summary['high_priority_gaps']:
        items.append(f"Address the {summary['high_priority_gaps']} high-priority individual gap{'s' if summary['high_priority_gaps'] != 1 else ''} "
                     'listed in the training projections first; these are the gaps most likely to affect performance.')
    coverage = summary['assessment_coverage']
    if coverage is None:
        items.append('No required skills have been rated yet. Run the self-assessment and supervisor ratings before relying on this report.')
    elif coverage < MIN_COVERAGE:
        items.append(f'Only {coverage}% of required skills have been rated. Complete the outstanding assessments '
                     f'to reach at least {MIN_COVERAGE}% before submitting this report.')
    if summary['staff_without_role']:
        items.append(f"Assign job roles to the {summary['staff_without_role']} officer{'s' if summary['staff_without_role'] != 1 else ''} "
                     'without one so their needs can be assessed.')
    if not items:
        items.append('All assessed officers meet the requirements of their job roles. Review the role profiles and reassess next period.')
    items.append('Review the training needs again after the training is delivered and in the next staff appraisal cycle.')
    return items


def filename(report, extension):
    scope = report['scope']['org_unit']['name'] if report['scope']['org_unit'] else report['institution']['name']
    slug = re.sub(r'[^A-Za-z0-9]+', '-', f"TNA {scope} {report['period']}").strip('-')
    return f'{slug}.{extension}'


# Shared table layouts, so the PDF and the workbook carry the same columns.

def _pct(value):
    return '—' if value is None else f'{value}%'


def _num(value):
    return '—' if value is None else value


def _level(value):
    return '—' if value is None else f'{value} {LEVEL_NAMES.get(value, "")}'.strip()


def summary_rows(report):
    s = report['summary']
    rows = [
        ('Officers in scope', s['staff']),
        ('Officers assessed', s['staff_assessed']),
        ('Officers with at least one gap', s['staff_with_gaps']),
        ('Officers to train', s['officers_to_train']),
        ('Assessment coverage (required skills rated)', _pct(s['assessment_coverage'])),
        ('Gap score (0 = all requirements met)', _num(s['gap_score'])),
        ('Skills with gaps', s['skills_with_gaps']),
        ('High-priority individual gaps', s['high_priority_gaps']),
    ]
    if report['scope']['type'] == 'institution':
        rows.append(('Officers without a job role', s['staff_without_role']))
    return rows


def organisational_table(report):
    header = ['Department', 'Officers', 'Assessed', 'With gaps', 'Coverage', 'Gap score']
    rows = [['    ' * u['depth'] + u['org_unit']['name'], u['staff'], u['staff_assessed'], u['staff_with_gaps'], _pct(u['assessment_coverage']), _num(u['gap_score'])]
            for u in report['organisational']]
    return header, rows


def occupational_table(report):
    header = ['Job role', 'Department', 'Officers', 'Assessed', 'With gaps', 'Gap score', 'Main training needs']
    rows = [[r['job_role']['title'], r['org_unit']['name'] if r['org_unit'] else '—', r['staff'], r['staff_assessed'],
             r['staff_with_gaps'], _num(r['gap_score']), ', '.join(r['priority_needs']) or '—']
            for r in report['occupational']]
    return header, rows


def needs_table(report):
    header = ['#', 'Skill', 'Type', 'Officers requiring', 'Assessed', 'Officers to train', 'High priority', 'Average gap', 'Priority']
    rows = [[i, n['skill']['name'], n['skill']['type'].title(), n['staff_requiring'], n['staff_assessed'], n['officers_to_train'],
             n['high_priority'], _num(n['average_gap']), PRIORITY_NAMES[n['priority']]]
            for i, n in enumerate(report['training_needs'], start=1)]
    return header, rows


def individuals_table(report):
    header = ['Officer', 'P/F No.', 'Job role', 'Department', 'Required', 'Assessed', 'Gap score', 'Skill gaps']
    rows = []
    for p in report['individuals']:
        gaps = ', '.join(f"{g['skill']['name']} (-{g['gap']})" for g in p['gaps']) or ('None' if p['skills_assessed'] else 'Not assessed')
        rows.append([_name(p['user']), p['user']['username'], p['job_role']['title'] if p['job_role'] else '—',
                     p['org_unit']['name'] if p['org_unit'] else '—', p['skills_required'], p['skills_assessed'], _num(p['gap_score']), gaps])
    return header, rows


def projections_table(report):
    header = ['#', 'Skill / training area', 'Officer', 'P/F No.', 'Job role', 'Department', 'Required', 'Current', 'Gap', 'Priority']
    rows = [[i, p['skill']['name'], _name(p['officer']), p['officer']['username'], p['job_role']['title'] if p['job_role'] else '—',
             p['org_unit']['name'] if p['org_unit'] else '—', _level(p['required_level']), _level(p['current_level']), p['gap'],
             PRIORITY_NAMES[p['priority']]]
            for i, p in enumerate(report['projections'], start=1)]
    return header, rows


def sections(report):
    """(number, title, intro, table) for each tabular section, in report order."""
    return [
        ('3', 'Organisational analysis', 'Gaps by department. Each line includes the units below it, shown indented.', organisational_table(report)),
        ('4', 'Occupational analysis', 'Gaps by job role, measured against each role\'s skill profile.', occupational_table(report)),
        ('5', 'Individual analysis', 'Each officer\'s gaps against their job role. P/F No. is the officer\'s LSMS username.', individuals_table(report)),
        ('6', 'Prioritised training needs', 'Skills ranked by the total DIF-weighted gap across all officers.', needs_table(report)),
    ]


def scope_label(report):
    if report['scope']['org_unit']:
        return f"{report['scope']['org_unit']['name']} and the units below it"
    return 'Whole institution'


# PDF

def render_pdf(report):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    body = ParagraphStyle('body', parent=styles['BodyText'], fontSize=9.5, leading=13)
    cell = ParagraphStyle('cell', parent=body, fontSize=8, leading=10)
    head_cell = ParagraphStyle('head', parent=cell, fontName='Helvetica-Bold', textColor=colors.white)
    h1 = ParagraphStyle('h1', parent=styles['Title'], fontSize=18, spaceAfter=4)
    h2 = ParagraphStyle('h2', parent=styles['Heading2'], fontSize=13, spaceBefore=10, spaceAfter=4)
    sub = ParagraphStyle('sub', parent=body, alignment=1, textColor=colors.HexColor('#444444'))
    navy = colors.HexColor('#1f3a5f')

    def esc(value):
        return str(value).replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')

    def table(header, rows, widths=None):
        if not rows:
            return Paragraph('<i>Nothing to report.</i>', body)
        data = [[Paragraph(esc(h), head_cell) for h in header]] + [[Paragraph(esc(v).replace('    ', '&nbsp;' * 4), cell) for v in row] for row in rows]
        t = Table(data, colWidths=widths, repeatRows=1, hAlign='LEFT')
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), navy),
            ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#b8c2cc')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f2f5f8')]),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('TOPPADDING', (0, 0), (-1, -1), 3),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ]))
        return t

    pagesize = landscape(A4)
    width = pagesize[0] - 30 * mm
    institution = report['institution']['name']

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont('Helvetica', 7.5)
        canvas.setFillColor(colors.HexColor('#666666'))
        canvas.drawString(15 * mm, 9 * mm, f"{institution} · Training Needs Assessment {report['period']}")
        canvas.drawRightString(pagesize[0] - 15 * mm, 9 * mm, f'Page {doc.page}')
        canvas.restoreState()

    s = report['summary']
    story = [
        Paragraph(esc(institution.upper()), sub),
        Spacer(1, 4),
        Paragraph(report['title'], h1),
        Paragraph(esc(f"Financial year {report['period']} · {scope_label(report)}"), sub),
        Paragraph(esc(f"Generated {report['generated_at'][:10]} from LSMS"), sub),
        Spacer(1, 8),
        Paragraph('1. Executive summary', h2),
        Paragraph(esc(
            f"{s['staff']} officers were in scope and {s['staff_assessed']} were assessed, covering "
            f"{_pct(s['assessment_coverage'])} of the skills their job roles require. {s['officers_to_train']} officers have at least one "
            f"gap across {s['skills_with_gaps']} skills, {s['high_priority_gaps']} of them high priority. The overall gap score is "
            f"{_num(s['gap_score'])} out of 100."), body),
        Spacer(1, 4),
        table(['Indicator', 'Value'], summary_rows(report), [width * 0.45, width * 0.2]),
        Paragraph('2. Methodology', h2),
        *[Paragraph(esc(f'• {line}'), body) for line in report['methodology']],
    ]
    widths = {
        '3': [0.34, 0.1, 0.1, 0.1, 0.12, 0.12],
        '4': [0.2, 0.17, 0.08, 0.08, 0.08, 0.08, 0.31],
        '5': [0.16, 0.1, 0.14, 0.13, 0.07, 0.07, 0.07, 0.26],
        '6': [0.04, 0.24, 0.09, 0.11, 0.09, 0.11, 0.1, 0.1, 0.1],
    }
    for number, title, intro, (header, rows) in sections(report):
        story += [KeepTogether([Paragraph(f'{number}. {title}', h2), Paragraph(esc(intro), body), Spacer(1, 3)]),
                  table(header, rows, [width * w for w in widths[number]])]
    story += [Paragraph('7. Recommendations', h2), *[Paragraph(esc(f'{i}. {line}'), body) for i, line in enumerate(report['recommendations'], 1)]]
    story += [
        Spacer(1, 14),
        table(['', 'Name', 'Signature', 'Date'], [['Prepared by (Head of HRM&D)', '', '', ''], ['Approved by (Authorised Officer)', '', '', '']],
              [width * 0.3, width * 0.3, width * 0.2, width * 0.15]),
        PageBreak(),
        Paragraph('Annex: Training projections', h2),
        Paragraph('Officers proposed for training in each skill, highest priority first, for the Human Resource Management '
                  'Advisory Committee to consider when nominating officers for training.', body),
        Spacer(1, 3),
        table(*projections_table(report), [width * w for w in (0.04, 0.17, 0.14, 0.09, 0.13, 0.12, 0.09, 0.09, 0.05, 0.08)]),
    ]

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=pagesize, leftMargin=15 * mm, rightMargin=15 * mm, topMargin=14 * mm, bottomMargin=16 * mm,
                            title=f"{report['title']} {report['period']}", author=institution)
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buffer.getvalue()


# Excel

def render_xlsx(report):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    bold = Font(bold=True)
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill('solid', fgColor='1F3A5F')
    wrap = Alignment(wrap_text=True, vertical='top')

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Summary'
    sheet.append([report['institution']['name']])
    sheet.append([report['title']])
    sheet['A1'].font = Font(bold=True, size=12)
    sheet['A2'].font = Font(bold=True, size=14)
    sheet.append(['Financial year', report['period']])
    sheet.append(['Scope', scope_label(report)])
    sheet.append(['Generated', report['generated_at']])
    sheet.append([])
    sheet.append(['Indicator', 'Value'])
    for c in sheet[sheet.max_row]:
        c.font, c.fill = header_font, header_fill
    for label, value in summary_rows(report):
        sheet.append([label, value])
    sheet.append([])
    sheet.append(['Methodology'])
    sheet.cell(sheet.max_row, 1).font = bold
    for line in report['methodology']:
        sheet.append([line])
    sheet.append([])
    sheet.append(['Recommendations'])
    sheet.cell(sheet.max_row, 1).font = bold
    for i, line in enumerate(report['recommendations'], 1):
        sheet.append([f'{i}. {line}'])
    sheet.column_dimensions['A'].width = 48
    sheet.column_dimensions['B'].width = 28

    def add_sheet(title, header, rows):
        ws = workbook.create_sheet(title)
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill, c.alignment = header_font, header_fill, wrap
        for row in rows:
            ws.append(row)
        ws.freeze_panes = 'A2'
        if rows:
            ws.auto_filter.ref = ws.dimensions
        for index, name in enumerate(header, 1):
            longest = max([len(str(name))] + [len(str(row[index - 1])) for row in rows])
            ws.column_dimensions[get_column_letter(index)].width = min(max(longest + 2, 8), 60)

    add_sheet('Departments', *organisational_table(report))
    add_sheet('Job roles', *occupational_table(report))
    add_sheet('Training needs', *needs_table(report))
    add_sheet('Individuals', *individuals_table(report))
    # Projections keep numeric levels so they can be sorted and summed.
    header = ['#', 'Skill / training area', 'Officer', 'P/F No.', 'Job role', 'Department', 'Required level', 'Current level',
              'Rated by', 'Gap', 'Priority score', 'Priority']
    rows = [[i, p['skill']['name'], _name(p['officer']), p['officer']['username'], p['job_role']['title'] if p['job_role'] else '',
             p['org_unit']['name'] if p['org_unit'] else '', p['required_level'], p['current_level'],
             'Supervisor' if p['rating_source'] == 'manager' else 'Self', p['gap'], p['priority_score'], PRIORITY_NAMES[p['priority']]]
            for i, p in enumerate(report['projections'], 1)]
    add_sheet('Training projections', header, rows)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
