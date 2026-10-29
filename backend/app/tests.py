import datetime
import io
from datetime import timedelta
from unittest import mock

from django.contrib.auth import get_user_model
from django.core import mail
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import TestCase, override_settings
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient

from .models import (
    Grade, JobFamily, CompetencyFramework, ImportRun,
    Tenant, OrgUnit, JobRole, Skill, JobRoleSkill, EmployeeSkillRating, TrainingCalendar,
    Questionnaire, Question, QuestionChoice, QuestionResponse, Invitation, TrainingPlan, TrainingNomination,
    License, LicenseEvent,
)
from . import licensing

User = get_user_model()
PASSWORD = 'a-Strong-pass-123'
FAST_HASHER = override_settings(PASSWORD_HASHERS=['django.contrib.auth.hashers.MD5PasswordHasher'])


def make_tenant_data(tenant, user):
    """Create one of every tenant-owned object for ``tenant``."""
    org_unit = OrgUnit.objects.create(tenant=tenant, name=f'{tenant.name} HQ')
    job_role = JobRole.objects.create(tenant=tenant, org_unit=org_unit, title=f'{tenant.name} Engineer')
    skill = Skill.objects.create(tenant=tenant, name=f'{tenant.name} Python', difficulty=3, usage_frequency=4, impact_if_missing=5)
    questionnaire = Questionnaire.objects.create(tenant=tenant, title=f'{tenant.name} Survey', org_unit=org_unit)
    question = Question.objects.create(questionnaire=questionnaire, text='How good are you?', skill=skill)
    return {
        'org-units': org_unit,
        'job-roles': job_role,
        'skills': skill,
        'job-role-skills': JobRoleSkill.objects.create(tenant=tenant, job_role=job_role, skill=skill),
        'ratings': EmployeeSkillRating.objects.create(tenant=tenant, user=user, skill=skill, self_rating=3, survey_date=datetime.date(2026, 1, 1)),
        'training-calendar': TrainingCalendar.objects.create(tenant=tenant, skill=skill, training_title='Intro', scheduled_date=datetime.date(2026, 2, 1), format='online', created_by=user),
        'questionnaires': questionnaire,
        'questions': question,
        'question-choices': QuestionChoice.objects.create(question=question, text='Very'),
        'question-responses': QuestionResponse.objects.create(question=question, user=user, response_rating=4),
    }


@FAST_HASHER
class TenantTestCase(TestCase):
    def setUp(self):
        self.tenant_a = Tenant.objects.create(name='Acme')
        self.tenant_b = Tenant.objects.create(name='Globex')
        self.admin_a = User.objects.create_user(username='admin_a', password=PASSWORD, tenant=self.tenant_a, role='admin')
        self.user_a = User.objects.create_user(username='user_a', password=PASSWORD, tenant=self.tenant_a)
        self.user_b = User.objects.create_user(username='user_b', password=PASSWORD, tenant=self.tenant_b)
        self.data_a = make_tenant_data(self.tenant_a, self.user_a)
        self.data_b = make_tenant_data(self.tenant_b, self.user_b)
        self.client = APIClient()

    def login(self, user):
        self.client.force_authenticate(user=user)

    def ids(self, response):
        self.assertEqual(response.status_code, 200, response.content)
        return {str(row['id']) for row in response.json()}


class AuthenticationRequiredTests(TenantTestCase):
    ENDPOINTS = [
        'skills', 'job-roles', 'job-role-skills', 'ratings', 'training-calendar', 'org-units',
        'tenants', 'questionnaires', 'questions', 'question-choices', 'question-responses', 'users', 'invitations',
        'grades', 'job-families', 'competency-frameworks', 'import-runs',
    ]

    def test_anonymous_cannot_list_anything(self):
        for endpoint in self.ENDPOINTS:
            with self.subTest(endpoint=endpoint):
                response = self.client.get(f'/api/{endpoint}/')
                self.assertEqual(response.status_code, 401)

    def test_anonymous_cannot_read_a_user(self):
        response = self.client.get(f'/api/users/{self.user_a.pk}/')
        self.assertEqual(response.status_code, 401)

    def test_anonymous_cannot_create_tenant_data(self):
        response = self.client.post('/api/org-units/', {'name': 'Sneaky'})
        self.assertEqual(response.status_code, 401)


class TenantIsolationTests(TenantTestCase):
    def test_lists_only_show_own_tenant(self):
        self.login(self.user_a)
        for endpoint, obj in self.data_a.items():
            with self.subTest(endpoint=endpoint):
                ids = self.ids(self.client.get(f'/api/{endpoint}/'))
                self.assertIn(str(obj.pk), ids)
                self.assertNotIn(str(self.data_b[endpoint].pk), ids)

    def test_other_tenants_objects_are_not_found(self):
        self.login(self.admin_a)
        for endpoint, obj in self.data_b.items():
            with self.subTest(endpoint=endpoint):
                self.assertEqual(self.client.get(f'/api/{endpoint}/{obj.pk}/').status_code, 404)
                self.assertEqual(self.client.patch(f'/api/{endpoint}/{obj.pk}/', {}).status_code, 404)
                self.assertEqual(self.client.delete(f'/api/{endpoint}/{obj.pk}/').status_code, 404)
        self.assertTrue(OrgUnit.objects.filter(pk=self.data_b['org-units'].pk).exists())

    def test_users_list_only_shows_own_tenant(self):
        self.login(self.user_a)
        ids = self.ids(self.client.get('/api/users/'))
        self.assertEqual(ids, {str(self.admin_a.pk), str(self.user_a.pk)})

    def test_tenants_list_only_shows_own_tenant(self):
        self.login(self.user_a)
        self.assertEqual(self.ids(self.client.get('/api/tenants/')), {str(self.tenant_a.pk)})
        self.assertEqual(self.client.get(f'/api/tenants/{self.tenant_b.pk}/').status_code, 404)

    def test_create_assigns_own_tenant_even_if_another_is_sent(self):
        self.login(self.admin_a)
        response = self.client.post('/api/org-units/', {'name': 'Ops', 'tenant': str(self.tenant_b.pk)})
        self.assertEqual(response.status_code, 201, response.content)
        self.assertEqual(OrgUnit.objects.get(pk=response.json()['id']).tenant, self.tenant_a)

    def test_update_cannot_move_object_to_another_tenant(self):
        self.login(self.admin_a)
        org_unit = self.data_a['org-units']
        response = self.client.patch(f'/api/org-units/{org_unit.pk}/', {'tenant': str(self.tenant_b.pk)})
        self.assertEqual(response.status_code, 200, response.content)
        org_unit.refresh_from_db()
        self.assertEqual(org_unit.tenant, self.tenant_a)

    def test_cannot_reference_another_tenants_objects(self):
        self.login(self.admin_a)
        response = self.client.post('/api/job-role-skills/', {
            'job_role': str(self.data_a['job-roles'].pk),
            'skill': str(self.data_b['skills'].pk),
        })
        self.assertEqual(response.status_code, 400)
        self.assertIn('skill', response.json())

    def test_cannot_add_question_to_another_tenants_questionnaire(self):
        self.login(self.admin_a)
        response = self.client.post('/api/questions/', {
            'questionnaire': str(self.data_b['questionnaires'].pk),
            'text': 'Injected?',
        })
        self.assertEqual(response.status_code, 400)
        self.assertFalse(Question.objects.filter(text='Injected?').exists())

    def test_cannot_add_choice_to_another_tenants_question(self):
        self.login(self.admin_a)
        response = self.client.post('/api/question-choices/', {
            'question': str(self.data_b['questions'].pk),
            'text': 'Injected',
        })
        self.assertEqual(response.status_code, 400)

    def test_response_is_recorded_for_the_requesting_user(self):
        self.login(self.user_a)
        response = self.client.post('/api/question-responses/', {
            'question': str(self.data_a['questions'].pk),
            'user': str(self.user_b.pk),
            'response_rating': 5,
        })
        self.assertEqual(response.status_code, 201, response.content)
        self.assertEqual(QuestionResponse.objects.get(pk=response.json()['id']).user, self.user_a)

    def test_cannot_respond_to_another_tenants_question(self):
        self.login(self.user_a)
        response = self.client.post('/api/question-responses/', {
            'question': str(self.data_b['questions'].pk),
            'response_rating': 5,
        })
        self.assertEqual(response.status_code, 400)


class StructurePermissionTests(TenantTestCase):
    def test_employees_can_read_but_not_change_structure(self):
        self.login(self.user_a)
        for endpoint in ('org-units', 'job-roles', 'skills', 'job-role-skills'):
            with self.subTest(endpoint=endpoint):
                obj = self.data_a[endpoint]
                self.assertEqual(self.client.get(f'/api/{endpoint}/{obj.pk}/').status_code, 200)
                self.assertEqual(self.client.post(f'/api/{endpoint}/', {}).status_code, 403)
                self.assertEqual(self.client.patch(f'/api/{endpoint}/{obj.pk}/', {}).status_code, 403)
                self.assertEqual(self.client.delete(f'/api/{endpoint}/{obj.pk}/').status_code, 403)

    def test_admin_builds_org_structure(self):
        self.login(self.admin_a)
        parent = self.client.post('/api/org-units/', {'name': 'Engineering'}).json()['id']
        child = self.client.post('/api/org-units/', {'name': 'Platform', 'parent': parent})
        self.assertEqual(child.status_code, 201, child.content)
        role = self.client.post('/api/job-roles/', {'title': 'SRE', 'org_unit': child.json()['id']})
        self.assertEqual(role.status_code, 201, role.content)
        skill = self.client.post('/api/skills/', {'name': 'Kubernetes', 'type': 'technical', 'difficulty': 4, 'usage_frequency': 5, 'impact_if_missing': 5})
        self.assertEqual(skill.status_code, 201, skill.content)


class OrgStructureTests(TenantTestCase):
    def setUp(self):
        super().setUp()
        hq = self.data_a['org-units']
        self.eng = OrgUnit.objects.create(tenant=self.tenant_a, name='Engineering', parent=hq)
        self.platform = OrgUnit.objects.create(tenant=self.tenant_a, name='Platform', parent=self.eng)
        JobRole.objects.create(tenant=self.tenant_a, org_unit=self.platform, title='SRE')
        JobRole.objects.create(tenant=self.tenant_a, title='Floating role')

    def test_admin_gets_nested_tree_of_own_tenant(self):
        self.login(self.admin_a)
        data = self.client.get('/api/org-structure/').json()
        self.assertEqual(data['tenant']['name'], 'Acme')
        [hq] = data['units']
        self.assertEqual((hq['name'], [r['title'] for r in hq['job_roles']]), ('Acme HQ', ['Acme Engineer']))
        [eng] = hq['children']
        [platform] = eng['children']
        self.assertEqual((eng['name'], platform['name'], platform['job_roles'][0]['title']), ('Engineering', 'Platform', 'SRE'))
        self.assertEqual([r['title'] for r in data['unassigned_job_roles']], ['Floating role'])
        self.assertNotIn('Globex', str(data))

    def test_non_admins_are_refused(self):
        self.login(self.user_a)
        self.assertEqual(self.client.get('/api/org-structure/').status_code, 403)
        self.assertEqual(APIClient().get('/api/org-structure/').status_code, 401)

    def test_admin_cannot_view_another_tenant(self):
        self.login(self.admin_a)
        data = self.client.get(f'/api/org-structure/?tenant={self.tenant_b.pk}').json()
        self.assertEqual(data['tenant']['name'], 'Acme')

    def test_super_admin_picks_a_tenant(self):
        self.login(User.objects.create_superuser(username='root', password=PASSWORD, email='root@example.com'))
        self.assertEqual(self.client.get('/api/org-structure/').status_code, 400)
        self.assertEqual(self.client.get('/api/org-structure/?tenant=not-a-uuid').status_code, 404)
        data = self.client.get(f'/api/org-structure/?tenant={self.tenant_b.pk}').json()
        self.assertEqual([u['name'] for u in data['units']], ['Globex HQ'])

    def test_parent_cycle_does_not_hide_units(self):
        loop_a = OrgUnit.objects.create(tenant=self.tenant_a, name='Loop A')
        loop_b = OrgUnit.objects.create(tenant=self.tenant_a, name='Loop B', parent=loop_a)
        OrgUnit.objects.filter(pk=loop_a.pk).update(parent=loop_b)
        self.login(self.admin_a)
        names = [u['name'] for u in self.client.get('/api/org-structure/').json()['units']]
        self.assertIn('Loop A', names)


class TenantEndpointTests(TenantTestCase):
    def setUp(self):
        super().setUp()
        self.superadmin = User.objects.create_superuser(username='root', password=PASSWORD, email='root@example.com')

    def test_employee_cannot_edit_tenant(self):
        self.login(self.user_a)
        response = self.client.patch(f'/api/tenants/{self.tenant_a.pk}/', {'name': 'Renamed'})
        self.assertEqual(response.status_code, 403)

    def test_admin_can_edit_own_tenant(self):
        self.login(self.admin_a)
        response = self.client.patch(f'/api/tenants/{self.tenant_a.pk}/', {'name': 'Acme Ltd'})
        self.assertEqual(response.status_code, 200, response.content)
        self.tenant_a.refresh_from_db()
        self.assertEqual(self.tenant_a.name, 'Acme Ltd')

    def test_tenant_admins_cannot_create_or_delete_tenants(self):
        self.login(self.admin_a)
        self.assertEqual(self.client.post('/api/tenants/', {'name': 'New'}).status_code, 403)
        self.assertEqual(self.client.delete(f'/api/tenants/{self.tenant_a.pk}/').status_code, 403)
        self.assertEqual(Tenant.objects.count(), 2)

    def test_super_admin_creates_tenant_and_invites_its_first_admin(self):
        self.login(self.superadmin)
        response = self.client.post('/api/tenants/', {'name': 'Initech', 'admin_email': 'boss@initech.com'})
        self.assertEqual(response.status_code, 201, response.content)
        tenant = Tenant.objects.get(name='Initech')
        invitation = Invitation.objects.get(tenant=tenant)
        self.assertEqual((invitation.email, invitation.role, invitation.invited_by), ('boss@initech.com', 'admin', self.superadmin))
        self.assertEqual(len(mail.outbox), 1)
        self.assertIn(invitation.token, mail.outbox[0].body)

    def test_super_admin_sees_and_deletes_any_tenant(self):
        self.login(self.superadmin)
        self.assertEqual(self.ids(self.client.get('/api/tenants/')), {str(self.tenant_a.pk), str(self.tenant_b.pk)})
        self.assertEqual(self.client.delete(f'/api/tenants/{self.tenant_b.pk}/').status_code, 204)


class InvitationTests(TenantTestCase):
    def invite(self, user, **payload):
        self.login(user)
        return self.client.post('/api/invitations/', {'email': 'new@acme.com', **payload})

    def accept(self, token, **overrides):
        payload = {'username': 'newbie', 'password': PASSWORD, 'first_name': 'New'}
        payload.update(overrides)
        return APIClient().post(f'/api/invitations/accept/{token}/', payload)

    def test_admin_invites_into_own_tenant_and_email_is_sent(self):
        response = self.invite(self.admin_a, role='manager', tenant=str(self.tenant_b.pk))
        self.assertEqual(response.status_code, 201, response.content)
        invitation = Invitation.objects.get(email='new@acme.com')
        self.assertEqual((invitation.tenant, invitation.role, invitation.invited_by), (self.tenant_a, 'manager', self.admin_a))
        self.assertEqual(response.json()['status'], 'pending')
        self.assertIn(invitation.token, response.json()['accept_url'])
        self.assertEqual(mail.outbox[0].to, ['new@acme.com'])

    def test_employee_cannot_invite(self):
        self.assertEqual(self.invite(self.user_a).status_code, 403)

    def test_anonymous_cannot_invite(self):
        self.assertEqual(APIClient().post('/api/invitations/', {'email': 'x@y.com'}).status_code, 401)

    def test_cannot_invite_an_existing_users_email(self):
        self.user_b.email = 'taken@globex.com'
        self.user_b.save()
        self.assertEqual(self.invite(self.admin_a, email='Taken@Globex.com').status_code, 400)

    def test_admins_only_see_own_tenants_invitations(self):
        Invitation.objects.create(tenant=self.tenant_b, email='b@globex.com')
        mine = Invitation.objects.create(tenant=self.tenant_a, email='a@acme.com')
        self.login(self.admin_a)
        self.assertEqual(self.ids(self.client.get('/api/invitations/')), {str(mine.pk)})

    def test_admin_can_revoke_an_invitation(self):
        invitation = Invitation.objects.create(tenant=self.tenant_a, email='a@acme.com')
        self.login(self.admin_a)
        self.assertEqual(self.client.delete(f'/api/invitations/{invitation.pk}/').status_code, 204)
        self.assertEqual(self.accept(invitation.token).status_code, 410)

    def test_preview_shows_who_is_inviting(self):
        invitation = Invitation.objects.create(tenant=self.tenant_a, email='a@acme.com', role='manager')
        response = APIClient().get(f'/api/invitations/accept/{invitation.token}/')
        self.assertEqual(response.json(), {'tenant_name': 'Acme', 'email': 'a@acme.com', 'role': 'manager'})
        self.assertEqual(APIClient().get('/api/invitations/accept/nope/').status_code, 404)

    def test_accepting_creates_user_in_the_invited_tenant(self):
        invitation = Invitation.objects.create(tenant=self.tenant_a, email='a@acme.com', role='manager')
        response = self.accept(invitation.token, tenant=str(self.tenant_b.pk), role='admin')
        self.assertEqual(response.status_code, 201, response.content)
        user = User.objects.get(username='newbie')
        self.assertEqual((user.tenant, user.role, user.email), (self.tenant_a, 'manager', 'a@acme.com'))
        token = APIClient().post('/api/token/', {'username': 'newbie', 'password': PASSWORD})
        self.assertEqual(token.status_code, 200, token.content)

    def test_invitation_works_only_once(self):
        invitation = Invitation.objects.create(tenant=self.tenant_a, email='a@acme.com')
        self.assertEqual(self.accept(invitation.token).status_code, 201)
        self.assertEqual(self.accept(invitation.token, username='second').status_code, 410)
        self.assertEqual(APIClient().get(f'/api/invitations/accept/{invitation.token}/').status_code, 410)

    def test_expired_invitation_is_refused(self):
        invitation = Invitation.objects.create(tenant=self.tenant_a, email='a@acme.com', expires_at=timezone.now() - timedelta(minutes=1))
        self.assertEqual(self.accept(invitation.token).status_code, 410)
        self.assertFalse(User.objects.filter(username='newbie').exists())

    def test_weak_password_and_taken_username_are_rejected(self):
        invitation = Invitation.objects.create(tenant=self.tenant_a, email='a@acme.com')
        self.assertIn('password', self.accept(invitation.token, password='123').json())
        self.assertIn('username', self.accept(invitation.token, username='USER_A').json())
        invitation.refresh_from_db()
        self.assertIsNone(invitation.accepted_at)

    def test_first_user_in_a_tenant_becomes_admin(self):
        empty = Tenant.objects.create(name='Initech')
        invitation = Invitation.objects.create(tenant=empty, email='first@initech.com', role='employee')
        self.assertEqual(self.accept(invitation.token).status_code, 201)
        self.assertEqual(User.objects.get(username='newbie').role, 'admin')

    def test_new_admin_can_invite_others(self):
        empty = Tenant.objects.create(name='Initech')
        invitation = Invitation.objects.create(tenant=empty, email='first@initech.com', role='admin')
        self.accept(invitation.token)
        response = self.invite(User.objects.get(username='newbie'), email='second@initech.com')
        self.assertEqual(response.status_code, 201, response.content)
        self.assertEqual(Invitation.objects.get(email='second@initech.com').tenant, empty)


class UserEndpointTests(TenantTestCase):
    def test_there_is_no_public_signup(self):
        response = APIClient().post('/api/users/', {
            'username': 'newbie', 'password': PASSWORD, 'tenant': str(self.tenant_a.pk), 'role': 'admin',
        })
        self.assertEqual(response.status_code, 401)
        self.assertFalse(User.objects.filter(username='newbie').exists())

    def test_employee_cannot_add_users(self):
        self.login(self.user_a)
        response = self.client.post('/api/users/', {'username': 'x', 'password': PASSWORD})
        self.assertEqual(response.status_code, 403)

    def test_admin_adds_users_to_own_tenant(self):
        self.login(self.admin_a)
        response = self.client.post('/api/users/', {'username': 'hire', 'password': PASSWORD, 'tenant': str(self.tenant_b.pk)})
        self.assertEqual(response.status_code, 201, response.content)
        self.assertEqual(User.objects.get(username='hire').tenant, self.tenant_a)

    def test_user_can_edit_self_but_not_role(self):
        self.login(self.user_a)
        response = self.client.patch(f'/api/users/{self.user_a.pk}/', {'first_name': 'Ann'})
        self.assertEqual(response.status_code, 200, response.content)
        response = self.client.patch(f'/api/users/{self.user_a.pk}/', {'role': 'admin'})
        self.assertEqual(response.status_code, 403)
        self.user_a.refresh_from_db()
        self.assertEqual((self.user_a.first_name, self.user_a.role), ('Ann', 'employee'))

    def test_user_cannot_move_self_to_another_tenant(self):
        self.login(self.user_a)
        self.client.patch(f'/api/users/{self.user_a.pk}/', {'tenant': str(self.tenant_b.pk)})
        self.user_a.refresh_from_db()
        self.assertEqual(self.user_a.tenant, self.tenant_a)

    def test_user_cannot_edit_or_delete_colleagues(self):
        self.login(self.user_a)
        self.assertEqual(self.client.patch(f'/api/users/{self.admin_a.pk}/', {'first_name': 'X'}).status_code, 403)
        self.assertEqual(self.client.delete(f'/api/users/{self.admin_a.pk}/').status_code, 403)

    def test_admin_can_promote_and_delete_users(self):
        self.login(self.admin_a)
        response = self.client.patch(f'/api/users/{self.user_a.pk}/', {'role': 'manager'})
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(self.client.delete(f'/api/users/{self.user_a.pk}/').status_code, 204)

    def test_password_change_is_hashed(self):
        self.login(self.user_a)
        self.client.patch(f'/api/users/{self.user_a.pk}/', {'password': 'another-Pass-456'})
        self.user_a.refresh_from_db()
        self.assertTrue(self.user_a.check_password('another-Pass-456'))

    def test_me_returns_the_logged_in_user(self):
        self.login(self.admin_a)
        data = self.client.get('/api/users/me/').json()
        self.assertEqual((data['username'], data['role'], data['tenant'], data['is_superuser']), ('admin_a', 'admin', str(self.tenant_a.pk), False))
        self.assertEqual(APIClient().get('/api/users/me/').status_code, 401)

    def test_admin_can_deactivate_users_but_employees_cannot(self):
        self.login(self.user_a)
        self.assertEqual(self.client.patch(f'/api/users/{self.user_a.pk}/', {'is_active': False}).status_code, 403)
        self.login(self.admin_a)
        self.assertEqual(self.client.patch(f'/api/users/{self.user_a.pk}/', {'is_active': False}).status_code, 200)
        self.user_a.refresh_from_db()
        self.assertFalse(self.user_a.is_active)

    def test_admin_cannot_lock_themselves_out(self):
        self.login(self.admin_a)
        self.assertEqual(self.client.patch(f'/api/users/{self.admin_a.pk}/', {'role': 'employee'}).status_code, 403)
        self.assertEqual(self.client.patch(f'/api/users/{self.admin_a.pk}/', {'is_active': False}).status_code, 403)
        self.assertEqual(self.client.delete(f'/api/users/{self.admin_a.pk}/').status_code, 403)
        self.assertEqual(self.client.patch(f'/api/users/{self.admin_a.pk}/', {'first_name': 'Ada'}).status_code, 200)


class SuperAdminTests(TenantTestCase):
    def setUp(self):
        super().setUp()
        self.superadmin = User.objects.create_superuser(username='root', password=PASSWORD, email='root@example.com')
        self.login(self.superadmin)

    def test_tenant_list_includes_counts(self):
        Invitation.objects.create(tenant=self.tenant_a, email='pending@acme.com')
        Invitation.objects.create(tenant=self.tenant_a, email='old@acme.com', expires_at=timezone.now() - timedelta(days=1))
        acme = next(t for t in self.client.get('/api/tenants/').json() if t['name'] == 'Acme')
        self.assertEqual(
            (acme['user_count'], acme['org_unit_count'], acme['job_role_count'], acme['skill_count'], acme['pending_invitation_count']),
            (2, 1, 1, 1, 1),
        )

    def test_reads_every_tenants_data_and_can_filter_by_tenant(self):
        everything = self.ids(self.client.get('/api/org-units/'))
        self.assertEqual(everything, {str(self.data_a['org-units'].pk), str(self.data_b['org-units'].pk)})
        only_b = self.ids(self.client.get(f'/api/users/?tenant={self.tenant_b.pk}'))
        self.assertEqual(only_b, {str(self.user_b.pk)})
        only_a = self.ids(self.client.get(f'/api/invitations/?tenant={self.tenant_a.pk}'))
        self.assertEqual(only_a, set())

    def test_cannot_change_organisation_data(self):
        org_unit = self.data_a['org-units']
        self.assertEqual(self.client.post('/api/org-units/', {'name': 'X'}).status_code, 403)
        self.assertEqual(self.client.patch(f'/api/org-units/{org_unit.pk}/', {'name': 'X'}).status_code, 403)
        self.assertEqual(self.client.delete(f'/api/org-units/{org_unit.pk}/').status_code, 403)

    def test_can_manage_users_in_any_tenant(self):
        response = self.client.patch(f'/api/users/{self.user_b.pk}/', {'role': 'admin'})
        self.assertEqual(response.status_code, 200, response.content)
        self.user_b.refresh_from_db()
        self.assertEqual(self.user_b.role, 'admin')

    def test_can_invite_into_any_tenant(self):
        response = self.client.post('/api/invitations/', {'email': 'new@globex.com', 'role': 'admin', 'tenant': str(self.tenant_b.pk)})
        self.assertEqual(response.status_code, 201, response.content)
        self.assertEqual(Invitation.objects.get(email='new@globex.com').tenant, self.tenant_b)

    def test_me_has_no_tenant(self):
        data = self.client.get('/api/users/me/').json()
        self.assertEqual((data['is_superuser'], data['tenant']), (True, None))


class ModelTests(TestCase):
    def test_importance_score_is_mean_of_skill_factors(self):
        tenant = Tenant.objects.create(name='Acme')
        skill = Skill.objects.create(tenant=tenant, name='SQL', difficulty=2, usage_frequency=4, impact_if_missing=6)
        job_role = JobRole.objects.create(tenant=tenant, title='Analyst')
        link = JobRoleSkill.objects.create(tenant=tenant, job_role=job_role, skill=skill)
        self.assertEqual(link.importance_score, 4)
        self.assertEqual(str(link), 'Analyst - SQL (Importance: 4.0)')

    def test_role_dif_overrides_the_skills_own_scores(self):
        tenant = Tenant.objects.create(name='Acme')
        skill = Skill.objects.create(tenant=tenant, name='SQL', difficulty=2, usage_frequency=4, impact_if_missing=3)
        job_role = JobRole.objects.create(tenant=tenant, title='Analyst')
        link = JobRoleSkill.objects.create(tenant=tenant, job_role=job_role, skill=skill, importance=5, frequency=1)
        self.assertEqual(link.dif, (2, 5, 1))
        self.assertEqual(link.importance_score, 8 / 3)


class TokenTests(TenantTestCase):
    def test_token_endpoints_are_public(self):
        response = self.client.post('/api/token/', {'username': 'user_a', 'password': PASSWORD})
        self.assertEqual(response.status_code, 200, response.content)
        refresh = self.client.post('/api/token/refresh/', {'refresh': response.json()['refresh']})
        self.assertEqual(refresh.status_code, 200, refresh.content)

    def test_jwt_grants_access(self):
        access = self.client.post('/api/token/', {'username': 'user_a', 'password': PASSWORD}).json()['access']
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {access}')
        self.assertEqual(self.ids(self.client.get('/api/skills/')), {str(self.data_a['skills'].pk)})


@FAST_HASHER
class MinistryTestCase(TestCase):
    """
    A small ministry: Finance sits under Headquarters, HR is its own top-level unit.

    Accountants need IFMIS at level 4 (DIF 4/5/5 for the role, 4.67) and report
    writing at level 3 (the skill's own DIF 2/2/3, 2.33). HR officers need report
    writing at level 4.
    """

    def setUp(self):
        self.tenant = Tenant.objects.create(name='Ministry')
        self.other_tenant = Tenant.objects.create(name='County')
        hq = OrgUnit.objects.create(tenant=self.tenant, name='Headquarters')
        self.finance = OrgUnit.objects.create(tenant=self.tenant, name='Finance', parent=hq)
        self.hq = hq
        hr = OrgUnit.objects.create(tenant=self.tenant, name='Human Resources')
        accountant = JobRole.objects.create(tenant=self.tenant, title='Accountant', org_unit=self.finance)
        hr_officer = JobRole.objects.create(tenant=self.tenant, title='HR Officer', org_unit=hr)
        self.ifmis = Skill.objects.create(tenant=self.tenant, name='IFMIS', difficulty=4, usage_frequency=4, impact_if_missing=3)
        self.writing = Skill.objects.create(tenant=self.tenant, name='Report writing', type='soft', difficulty=2, usage_frequency=3, impact_if_missing=2)
        JobRoleSkill.objects.create(tenant=self.tenant, job_role=accountant, skill=self.ifmis, required_level=4, importance=5, frequency=5)
        JobRoleSkill.objects.create(tenant=self.tenant, job_role=accountant, skill=self.writing, required_level=3)
        JobRoleSkill.objects.create(tenant=self.tenant, job_role=hr_officer, skill=self.writing, required_level=4)

        def person(username, job_role=None, role='employee'):
            return User.objects.create_user(username=username, password=PASSWORD, tenant=self.tenant, role=role, job_role=job_role)

        self.alice = person('alice', accountant)
        self.bob = person('bob', accountant)
        self.carol = person('carol', hr_officer)
        self.admin = person('admin', role='admin')
        self.manager = person('manager', role='manager')
        self.outsider = User.objects.create_user(username='outsider', password=PASSWORD, tenant=self.other_tenant, role='admin')
        self.superadmin = User.objects.create_superuser(username='root', password=PASSWORD)

        def rate(user, skill, self_rating, manager_rating=None, when=datetime.date(2026, 6, 1)):
            EmployeeSkillRating.objects.create(tenant=self.tenant, user=user, skill=skill, self_rating=self_rating, manager_rating=manager_rating, survey_date=when)

        rate(self.alice, self.ifmis, 4, manager_rating=2)
        rate(self.alice, self.writing, 3)
        rate(self.bob, self.ifmis, 1, when=datetime.date(2025, 6, 1))
        rate(self.bob, self.ifmis, 5)
        rate(self.carol, self.writing, 2)
        self.client = APIClient()

    def get(self, user, url):
        self.client.force_authenticate(user=user)
        return self.client.get(url)

    def ok(self, user, url):
        response = self.get(user, url)
        self.assertEqual(response.status_code, 200, response.content)
        return response.json()


class GapEngineTests(MinistryTestCase):
    def test_person_gaps_prefer_manager_rating_and_weight_by_role_dif(self):
        data = self.ok(self.alice, '/api/gaps/me/')
        self.assertEqual((data['job_role']['title'], data['org_unit']['name']), ('Accountant', 'Finance'))
        ifmis, writing = data['skills']
        self.assertEqual(ifmis['skill']['name'], 'IFMIS')
        self.assertEqual((ifmis['required_level'], ifmis['current_level'], ifmis['rating_source'], ifmis['gap']), (4, 2, 'manager', 2))
        self.assertEqual(ifmis['dif'], {'difficulty': 4, 'importance': 5, 'frequency': 5})
        self.assertEqual((ifmis['dif_score'], ifmis['priority_score'], ifmis['priority']), (4.67, 9.34, 'high'))
        self.assertEqual((writing['current_level'], writing['rating_source'], writing['gap'], writing['priority']), (3, 'self', 0, 'none'))
        # 4.67 x 2 short, out of 4.67 x 4 + 2.33 x 3 required.
        self.assertEqual(data['gap_score'], 36.4)
        self.assertEqual((data['skills_required'], data['skills_assessed'], data['skills_with_gap'], data['high_priority_gaps']), (2, 2, 1, 1))

    def test_latest_rating_counts_and_unrated_skills_are_not_assessed(self):
        data = self.ok(self.bob, f'/api/gaps/people/{self.bob.pk}/')
        ifmis, writing = data['skills']
        self.assertEqual((ifmis['current_level'], ifmis['gap']), (5, 0))
        self.assertEqual((writing['current_level'], writing['gap'], writing['priority']), (None, None, 'not_assessed'))
        self.assertEqual((data['gap_score'], data['skills_assessed']), (0.0, 1))

    def test_person_without_job_role_has_nothing_to_measure(self):
        data = self.ok(self.admin, '/api/gaps/me/')
        self.assertEqual((data['job_role'], data['gap_score'], data['skills']), (None, None, []))

    def test_department_rolls_up_its_staff_and_ranks_training_needs(self):
        data = self.ok(self.manager, f'/api/gaps/departments/{self.finance.pk}/')
        self.assertEqual(data['org_unit']['name'], 'Finance')
        self.assertEqual(data['parent']['name'], 'Headquarters')
        self.assertEqual((data['staff'], data['staff_assessed'], data['staff_with_gaps']), (2, 2, 1))
        self.assertEqual(data['assessment_coverage'], 75.0)
        ifmis = data['skills'][0]
        self.assertEqual(ifmis['skill']['name'], 'IFMIS')
        self.assertEqual(
            (ifmis['staff_requiring'], ifmis['staff_assessed'], ifmis['staff_with_gap'], ifmis['high_priority'], ifmis['average_gap'], ifmis['priority_total']),
            (2, 2, 1, 1, 1.0, 9.34),
        )
        self.assertEqual([p['user']['username'] for p in data['people']], ['alice', 'bob'])

    def test_department_includes_units_below_it(self):
        data = self.ok(self.admin, f'/api/gaps/departments/{self.hq.pk}/')
        self.assertEqual(data['staff'], 2)
        self.assertEqual([(u['org_unit']['name'], u['staff']) for u in data['units']], [('Finance', 2)])

    def test_institution_covers_everyone_with_a_breakdown_by_top_level_unit(self):
        data = self.ok(self.admin, '/api/gaps/institution/')
        self.assertEqual(data['tenant']['name'], 'Ministry')
        self.assertEqual((data['staff'], data['staff_without_role'], data['staff_with_gaps']), (5, 2, 2))
        self.assertEqual([u['org_unit']['name'] for u in data['units']], ['Headquarters', 'Human Resources'])
        hr = data['units'][1]
        self.assertEqual((hr['staff'], hr['gap_score']), (1, 50.0))
        self.assertEqual([s['skill']['name'] for s in data['skills']], ['IFMIS', 'Report writing'])
        self.assertEqual(data['people'][0]['user']['username'], 'carol')

    def test_employees_only_see_their_own_gaps(self):
        self.assertEqual(self.get(self.alice, f'/api/gaps/people/{self.alice.pk}/').status_code, 200)
        self.assertEqual(self.get(self.alice, f'/api/gaps/people/{self.carol.pk}/').status_code, 403)
        self.assertEqual(self.get(self.alice, f'/api/gaps/departments/{self.finance.pk}/').status_code, 403)
        self.assertEqual(self.get(self.alice, '/api/gaps/institution/').status_code, 403)
        self.assertEqual(self.ok(self.manager, f'/api/gaps/people/{self.carol.pk}/')['user']['username'], 'carol')

    def test_other_tenants_cannot_see_gaps(self):
        self.assertEqual(self.get(self.outsider, f'/api/gaps/people/{self.alice.pk}/').status_code, 404)
        self.assertEqual(self.get(self.outsider, f'/api/gaps/departments/{self.finance.pk}/').status_code, 404)
        self.assertEqual(self.ok(self.outsider, '/api/gaps/institution/')['staff'], 1)

    def test_anonymous_cannot_see_gaps(self):
        for url in ('/api/gaps/me/', f'/api/gaps/people/{self.alice.pk}/', f'/api/gaps/departments/{self.finance.pk}/', '/api/gaps/institution/'):
            with self.subTest(url=url):
                self.assertEqual(self.client.get(url).status_code, 401)

    def test_super_admin_picks_a_tenant(self):
        self.assertEqual(self.get(self.superadmin, '/api/gaps/institution/').status_code, 400)
        self.assertEqual(self.get(self.superadmin, '/api/gaps/institution/?tenant=not-a-uuid').status_code, 404)
        self.assertEqual(self.ok(self.superadmin, f'/api/gaps/institution/?tenant={self.tenant.pk}')['staff'], 5)
        self.assertEqual(self.ok(self.superadmin, f'/api/gaps/departments/{self.finance.pk}/')['staff'], 2)
        self.assertEqual(self.get(self.superadmin, '/api/gaps/me/').status_code, 404)

    def test_only_admins_assign_job_roles(self):
        hr_officer = self.carol.job_role
        self.client.force_authenticate(user=self.alice)
        response = self.client.patch(f'/api/users/{self.alice.pk}/', {'job_role': str(hr_officer.pk)})
        self.assertEqual(response.status_code, 403)
        self.client.force_authenticate(user=self.admin)
        response = self.client.patch(f'/api/users/{self.alice.pk}/', {'job_role': str(hr_officer.pk)})
        self.assertEqual(response.status_code, 200, response.content)
        self.alice.refresh_from_db()
        self.assertEqual(self.alice.job_role, hr_officer)

    def test_role_profile_validation(self):
        self.client.force_authenticate(user=self.admin)
        accountant = self.alice.job_role
        duplicate = self.client.post('/api/job-role-skills/', {'job_role': str(accountant.pk), 'skill': str(self.ifmis.pk)})
        self.assertEqual(duplicate.status_code, 400)
        self.assertIn('skill', duplicate.json())
        out_of_range = self.client.post('/api/job-role-skills/', {'job_role': str(self.carol.job_role.pk), 'skill': str(self.ifmis.pk), 'required_level': 6})
        self.assertEqual(out_of_range.status_code, 400)
        self.assertIn('required_level', out_of_range.json())
        created = self.client.post('/api/job-role-skills/', {'job_role': str(self.carol.job_role.pk), 'skill': str(self.ifmis.pk), 'required_level': 2, 'importance': 1})
        self.assertEqual(created.status_code, 201, created.content)
        # Difficulty and frequency fall back to the skill's own scores: (4 + 1 + 4) / 3.
        self.assertEqual(created.json()['importance_score'], 3)


class TNAReportTests(MinistryTestCase):
    URL = '/api/reports/tna/'

    def test_institution_report_follows_the_tna_structure(self):
        report = self.ok(self.admin, f'{self.URL}?period=2026/27')
        self.assertEqual((report['institution']['name'], report['period'], report['scope']['type']), ('Ministry', '2026/27', 'institution'))
        summary = report['summary']
        self.assertEqual((summary['staff'], summary['staff_without_role'], summary['officers_to_train'], summary['high_priority_gaps']), (5, 2, 2, 1))
        self.assertEqual(
            [(u['org_unit']['name'], u['depth'], u['staff']) for u in report['organisational']],
            [('Headquarters', 0, 2), ('Finance', 1, 2), ('Human Resources', 0, 1)],
        )
        self.assertEqual([r['job_role']['title'] for r in report['occupational']], ['HR Officer', 'Accountant'])
        self.assertEqual(report['occupational'][1]['priority_needs'], ['IFMIS'])
        self.assertEqual([(n['skill']['name'], n['officers_to_train']) for n in report['training_needs']], [('IFMIS', 1), ('Report writing', 1)])
        # Alice's IFMIS gap (priority 9.34) ranks above Carol's report writing gap (2 x 2.33).
        self.assertEqual(
            [(p['skill']['name'], p['officer']['username'], p['gap'], p['priority']) for p in report['projections']],
            [('IFMIS', 'alice', 2, 'high'), ('Report writing', 'carol', 2, 'medium')],
        )
        bob = next(p for p in report['individuals'] if p['user']['username'] == 'bob')
        self.assertEqual((bob['gaps'], bob['not_assessed']), ([], ['Report writing']))
        self.assertTrue(any('job roles' in line for line in report['recommendations']))

    def test_department_report_covers_its_subtree(self):
        report = self.ok(self.manager, f'{self.URL}?unit={self.hq.pk}')
        self.assertEqual((report['scope']['type'], report['scope']['org_unit']['name']), ('department', 'Headquarters'))
        self.assertEqual(report['summary']['staff'], 2)
        self.assertEqual([p['officer']['username'] for p in report['projections']], ['alice'])

    def test_period_defaults_to_the_current_financial_year(self):
        from .tna_report import financial_year
        self.assertEqual(financial_year(datetime.date(2026, 6, 30)), '2025/26')
        self.assertEqual(financial_year(datetime.date(2026, 7, 1)), '2026/27')
        self.assertRegex(self.ok(self.admin, self.URL)['period'], r'^\d{4}/\d{2}$')
        self.assertEqual(self.get(self.admin, f'{self.URL}?period=next-year').status_code, 400)

    def test_downloads_pdf_and_excel(self):
        from io import BytesIO
        from openpyxl import load_workbook

        pdf = self.get(self.admin, f'{self.URL}?period=2026/27&download=pdf')
        self.assertEqual(pdf.status_code, 200)
        self.assertEqual(pdf['Content-Type'], 'application/pdf')
        self.assertIn('TNA-Ministry-2026-27.pdf', pdf['Content-Disposition'])
        self.assertTrue(pdf.content.startswith(b'%PDF'))

        xlsx = self.get(self.admin, f'{self.URL}?unit={self.finance.pk}&period=2026/27&download=xlsx')
        self.assertEqual(xlsx.status_code, 200)
        self.assertIn('TNA-Finance-2026-27.xlsx', xlsx['Content-Disposition'])
        workbook = load_workbook(BytesIO(xlsx.content))
        self.assertEqual(workbook.sheetnames, ['Summary', 'Departments', 'Job roles', 'Training needs', 'Individuals', 'Training projections'])
        projections = list(workbook['Training projections'].values)
        self.assertEqual(projections[1][1:4], ('IFMIS', 'alice', 'alice'))
        self.assertEqual(self.get(self.admin, f'{self.URL}?download=docx').status_code, 400)

    def test_only_reviewers_in_the_tenant_get_the_report(self):
        self.assertEqual(self.get(self.alice, self.URL).status_code, 403)
        self.assertEqual(self.client.logout() or self.client.get(self.URL).status_code, 401)
        self.assertEqual(self.get(self.outsider, f'{self.URL}?unit={self.finance.pk}').status_code, 404)
        self.assertEqual(self.ok(self.outsider, self.URL)['summary']['staff'], 1)
        self.assertEqual(self.get(self.superadmin, self.URL).status_code, 400)
        self.assertEqual(self.ok(self.superadmin, f'{self.URL}?tenant={self.tenant.pk}')['summary']['staff'], 5)


class QuestionnaireAndRatingTests(TenantTestCase):
    def setUp(self):
        super().setUp()
        self.manager_a = User.objects.create_user(username='manager_a', password=PASSWORD, tenant=self.tenant_a, role='manager')
        self.skill = self.data_a['skills']
        self.questionnaire = self.data_a['questionnaires']
        self.question = self.data_a['questions']

    def test_only_admins_write_questionnaires(self):
        self.login(self.user_a)
        self.assertEqual(self.client.get('/api/questionnaires/').status_code, 200)
        self.assertEqual(self.client.post('/api/questionnaires/', {'title': 'Mine'}).status_code, 403)
        self.assertEqual(self.client.post('/api/questions/', {'questionnaire': str(self.questionnaire.pk), 'text': 'Q?'}).status_code, 403)
        self.login(self.admin_a)
        response = self.client.post('/api/questionnaires/', {'title': 'Annual TNA'})
        self.assertEqual(response.status_code, 201, response.content)
        response = self.client.post('/api/questions/', {'questionnaire': response.json()['id'], 'text': 'Rate your Python', 'skill': str(self.skill.pk)})
        self.assertEqual(response.status_code, 201, response.content)

    def test_lists_filter_by_query_parameters(self):
        other = Questionnaire.objects.create(tenant=self.tenant_a, title='Other')
        Question.objects.create(questionnaire=other, text='Elsewhere?')
        self.login(self.user_a)
        self.assertEqual(self.ids(self.client.get('/api/questions/', {'questionnaire': str(self.questionnaire.pk)})), {str(self.question.pk)})
        self.assertEqual(self.ids(self.client.get('/api/questions/', {'questionnaire': 'not-a-uuid'})), set())
        role = self.data_a['job-roles']
        self.assertEqual(self.ids(self.client.get('/api/job-role-skills/', {'job_role': str(role.pk)})), {str(self.data_a['job-role-skills'].pk)})

    def test_rating_answer_linked_to_a_skill_records_a_self_rating(self):
        self.login(self.user_a)
        response = self.client.post('/api/question-responses/', {'question': str(self.question.pk), 'response_rating': 2})
        self.assertEqual(response.status_code, 201, response.content)
        today = timezone.localdate()
        rating = EmployeeSkillRating.objects.get(user=self.user_a, skill=self.skill, survey_date=today)
        self.assertEqual(rating.self_rating, 2)
        # Answering again the same day updates that rating rather than adding another.
        self.client.post('/api/question-responses/', {'question': str(self.question.pk), 'response_rating': 4})
        rating.refresh_from_db()
        self.assertEqual(rating.self_rating, 4)
        self.assertEqual(EmployeeSkillRating.objects.filter(user=self.user_a, survey_date=today).count(), 1)

    def test_rating_answers_must_be_1_to_5(self):
        self.login(self.user_a)
        self.assertEqual(self.client.post('/api/question-responses/', {'question': str(self.question.pk), 'response_rating': 9}).status_code, 400)
        self.assertEqual(self.client.post('/api/question-responses/', {'question': str(self.question.pk)}).status_code, 400)

    def test_employees_only_see_their_own_answers_and_ratings(self):
        EmployeeSkillRating.objects.create(tenant=self.tenant_a, user=self.admin_a, skill=self.skill, self_rating=5, survey_date=datetime.date(2026, 1, 1))
        QuestionResponse.objects.create(question=self.question, user=self.admin_a, response_rating=5)
        self.login(self.user_a)
        self.assertEqual(self.ids(self.client.get('/api/ratings/')), {str(self.data_a['ratings'].pk)})
        self.assertEqual(self.ids(self.client.get('/api/question-responses/')), {str(self.data_a['question-responses'].pk)})
        self.login(self.manager_a)
        self.assertEqual(len(self.ids(self.client.get('/api/ratings/'))), 2)

    def test_people_rate_themselves_but_not_their_manager_rating(self):
        self.login(self.user_a)
        response = self.client.post('/api/ratings/', {'skill': str(self.skill.pk), 'user': str(self.admin_a.pk), 'self_rating': 3, 'survey_date': '2026-03-01'})
        self.assertEqual(response.status_code, 201, response.content)
        self.assertEqual(EmployeeSkillRating.objects.get(pk=response.json()['id']).user, self.user_a)
        response = self.client.post('/api/ratings/', {'skill': str(self.skill.pk), 'self_rating': 3, 'manager_rating': 5, 'survey_date': '2026-03-01'})
        self.assertEqual(response.status_code, 403)
        rating = self.data_a['ratings']
        self.assertEqual(self.client.patch(f'/api/ratings/{rating.pk}/', {'manager_rating': 5}).status_code, 403)
        self.assertEqual(self.client.patch(f'/api/ratings/{rating.pk}/', {'self_rating': 4}).status_code, 200)

    def test_managers_add_manager_ratings_to_others(self):
        rating = self.data_a['ratings']
        self.login(self.manager_a)
        self.assertEqual(self.client.patch(f'/api/ratings/{rating.pk}/', {'self_rating': 1}).status_code, 403)
        response = self.client.patch(f'/api/ratings/{rating.pk}/', {'manager_rating': 2})
        self.assertEqual(response.status_code, 200, response.content)
        rating.refresh_from_db()
        self.assertEqual((rating.self_rating, rating.manager_rating), (3, 2))
        self.assertEqual(self.client.delete(f'/api/ratings/{rating.pk}/').status_code, 403)


class TrainingPlanTests(MinistryTestCase):
    """
    From the gaps: Alice's IFMIS gap (2 levels, priority 9.34) and Carol's report
    writing gap (2 levels, 4.66). Each officer costs KES 20,000 in fees plus
    KES 5,000 a day for 5 days, so KES 45,000.
    """

    def setUp(self):
        super().setUp()
        self.plan = TrainingPlan.objects.create(
            tenant=self.tenant, period='2026/27', budget=80000, default_fee=20000, default_allowance_per_day=5000, default_duration_days=5,
        )

    def send(self, user, method, url, data=None):
        self.client.force_authenticate(user=user)
        return getattr(self.client, method)(url, data, format='json')

    def generate(self):
        response = self.send(self.admin, 'post', f'/api/training-plans/{self.plan.pk}/generate/')
        self.assertEqual(response.status_code, 200, response.content)
        return response.json()

    def nomination(self, username):
        return TrainingNomination.objects.get(programme__plan=self.plan, user__username=username)

    def approve_and_attend(self, username, attendance='attended'):
        nomination = self.nomination(username)
        response = self.send(self.admin, 'patch', f'/api/training-nominations/{nomination.pk}/', {'status': 'approved', 'attendance': attendance})
        self.assertEqual(response.status_code, 200, response.content)
        return nomination

    def test_generates_programmes_and_nominations_from_the_gaps(self):
        data = self.generate()
        self.assertEqual(data['created'], {'programmes': 2, 'nominations': 2})
        # Priority bands match the TNA report's training needs: total priority per officer assessed.
        self.assertEqual(
            [(p['title'], p['priority'], p['quarter'], p['participants'], p['cost_per_participant']) for p in data['programmes']],
            [('IFMIS', 'medium', 2, 1, 45000.0), ('Report writing', 'low', 3, 1, 45000.0)],
        )
        alice = self.nomination('alice')
        self.assertEqual((alice.source, alice.status, alice.required_level, alice.current_level, alice.gap, alice.priority_score),
                         ('gap', 'proposed', 4, 2, 2, 9.34))
        self.assertEqual(alice.nominated_by, self.admin)

        # Running it again adds only new gaps: Bob's first report writing rating shows a gap.
        self.assertEqual(self.generate()['created'], {'programmes': 0, 'nominations': 0})
        EmployeeSkillRating.objects.create(tenant=self.tenant, user=self.bob, skill=self.writing, self_rating=1, survey_date=datetime.date(2026, 7, 1))
        self.assertEqual(self.generate()['created'], {'programmes': 0, 'nominations': 1})

    def test_plan_is_costed_by_quarter_department_and_budget(self):
        self.generate()
        ifmis = self.plan.programmes.get(skill=self.ifmis)
        ifmis.fixed_cost = 10000
        ifmis.save()
        data = self.ok(self.manager, f'/api/training-plans/{self.plan.pk}/')
        self.assertEqual(data['summary'], {
            'programmes': 2, 'participants': 2, 'officers': 2, 'approved_participants': 0,
            'total_cost': 100000.0, 'budget': 80000.0, 'balance': -20000.0,
        })
        self.assertEqual([(q['quarter'], q['cost']) for q in data['by_quarter']], [(1, 0), (2, 55000.0), (3, 45000.0), (4, 0)])
        self.assertEqual([(d['org_unit']['name'], d['cost']) for d in data['by_department']], [('Finance', 55000.0), ('Human Resources', 45000.0)])

        # Declined nominees drop out of the cost; the fixed cost stays.
        self.send(self.admin, 'patch', f"/api/training-nominations/{self.nomination('alice').pk}/", {'status': 'declined'})
        data = self.ok(self.admin, f'/api/training-plans/{self.plan.pk}/')
        self.assertEqual((data['summary']['participants'], data['summary']['total_cost']), (1, 55000.0))
        listed = self.ok(self.admin, '/api/training-plans/')
        self.assertEqual(listed[0]['summary']['total_cost'], 55000.0)

    def test_attendance_and_evaluation_close_the_loop(self):
        self.generate()
        alice = self.nomination('alice')
        url = f'/api/training-nominations/{alice.pk}/evaluation/'
        self.assertEqual(self.send(self.alice, 'post', url, {'reaction_rating': 5}).status_code, 400)

        self.approve_and_attend('alice')
        self.assertEqual(self.send(self.alice, 'post', url, {'post_training_level': 5}).status_code, 403)
        response = self.send(self.alice, 'post', url, {'reaction_rating': 5, 'relevance_rating': 4, 'participant_comments': 'Useful'})
        self.assertEqual(response.status_code, 200, response.content)
        self.assertIsNotNone(response.json()['participant_submitted_at'])
        self.assertEqual(self.send(self.manager, 'post', url, {'reaction_rating': 1}).status_code, 403)

        response = self.send(self.manager, 'post', url, {'pre_test_score': 40, 'post_test_score': 75, 'post_training_level': 4})
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual((response.json()['reaction_rating'], response.json()['supervisor_name']), (5, 'manager'))

        # The supervisor's rating becomes Alice's manager rating, so the gap engine sees the gap closed.
        ifmis = self.ok(self.alice, '/api/gaps/me/')['skills'][0]
        self.assertEqual((ifmis['current_level'], ifmis['rating_source'], ifmis['gap']), (4, 'manager', 0))

        data = self.ok(self.admin, f'/api/training-plans/{self.plan.pk}/')
        self.assertEqual(data['attendance']['attended'], 1)
        self.assertEqual(data['attendance']['attendance_rate'], 100.0)
        self.assertEqual(data['evaluation'], {
            'participant_responses': 1, 'average_reaction': 5.0, 'average_relevance': 4.0, 'tested': 1,
            'average_learning_gain': 35.0, 'supervisor_evaluations': 1, 'gaps_closed': 1,
        })

    def test_who_can_do_what(self):
        self.generate()
        plans = '/api/training-plans/'
        self.assertEqual(self.get(self.alice, plans).status_code, 403)
        self.assertEqual(len(self.ok(self.manager, plans)), 1)
        self.assertEqual(self.send(self.manager, 'post', plans, {'period': '2027/28'}).status_code, 403)
        self.assertEqual(self.send(self.manager, 'post', f'/api/training-plans/{self.plan.pk}/generate/').status_code, 403)

        ifmis = self.plan.programmes.get(skill=self.ifmis)
        # Managers propose; HR approves.
        approved = self.send(self.manager, 'post', '/api/training-nominations/', {'programme': str(ifmis.pk), 'user': str(self.bob.pk), 'status': 'approved'})
        self.assertEqual(approved.status_code, 403)
        proposed = self.send(self.manager, 'post', '/api/training-nominations/', {'programme': str(ifmis.pk), 'user': str(self.bob.pk)})
        self.assertEqual(proposed.status_code, 201, proposed.content)
        self.assertEqual((proposed.json()['source'], proposed.json()['status']), ('manual', 'proposed'))
        duplicate = self.send(self.admin, 'post', '/api/training-nominations/', {'programme': str(ifmis.pk), 'user': str(self.bob.pk)})
        self.assertEqual(duplicate.status_code, 400)

        alice = self.nomination('alice')
        alice_url = f'/api/training-nominations/{alice.pk}/'
        self.assertEqual(self.send(self.manager, 'patch', alice_url, {'status': 'approved'}).status_code, 403)
        self.assertEqual(self.send(self.admin, 'patch', alice_url, {'attendance': 'attended'}).status_code, 400)
        self.assertEqual(self.send(self.admin, 'patch', alice_url, {'status': 'approved'}).status_code, 200)
        self.assertEqual(self.send(self.manager, 'patch', alice_url, {'attendance': 'partial', 'days_attended': 3}).status_code, 200)
        self.assertEqual(self.send(self.manager, 'patch', alice_url, {'days_attended': 9}).status_code, 400)
        self.assertEqual(self.send(self.alice, 'patch', alice_url, {'attendance': 'attended'}).status_code, 403)
        self.assertEqual(self.send(self.manager, 'delete', alice_url).status_code, 403)

        # Officers only see their own nominations; other tenants see nothing.
        self.assertEqual([n['officer']['username'] for n in self.ok(self.alice, '/api/training-nominations/')], ['alice'])
        self.assertEqual(len(self.ok(self.manager, f'/api/training-nominations/?plan={self.plan.pk}')), 3)
        self.assertEqual(self.get(self.outsider, f'/api/training-plans/{self.plan.pk}/').status_code, 404)
        self.assertEqual(self.ok(self.outsider, '/api/training-nominations/'), [])
        foreign = self.send(self.outsider, 'post', '/api/training-programmes/', {'plan': str(self.plan.pk), 'title': 'Sneaky'})
        self.assertEqual(foreign.status_code, 400)

    def test_period_status_and_approval(self):
        plans = '/api/training-plans/'
        self.assertEqual(self.send(self.admin, 'post', plans, {'period': '2026/27'}).status_code, 400)
        self.assertEqual(self.send(self.admin, 'post', plans, {'period': '2027/29'}).status_code, 400)
        created = self.send(self.admin, 'post', plans, {'period': '2027/28', 'budget': '1500000'})
        self.assertEqual(created.status_code, 201, created.content)
        self.assertEqual(created.json()['created_by'], str(self.admin.pk))

        url = f'{plans}{self.plan.pk}/'
        approved = self.send(self.admin, 'patch', url, {'status': 'approved'})
        self.assertEqual((approved.json()['status'], approved.json()['approved_by']), ('approved', str(self.admin.pk)))
        self.assertEqual(self.send(self.admin, 'post', f'{url}generate/').status_code, 400)
        self.assertIsNone(self.send(self.admin, 'patch', url, {'status': 'draft'}).json()['approved_at'])

    def test_downloads_pdf_and_excel(self):
        from io import BytesIO
        from openpyxl import load_workbook

        self.generate()
        self.approve_and_attend('alice')
        pdf = self.get(self.admin, f'/api/training-plans/{self.plan.pk}/export/?download=pdf')
        self.assertEqual(pdf.status_code, 200)
        self.assertIn('Training-plan-Ministry-2026-27.pdf', pdf['Content-Disposition'])
        self.assertTrue(pdf.content.startswith(b'%PDF'))

        xlsx = self.get(self.manager, f'/api/training-plans/{self.plan.pk}/export/?download=xlsx')
        self.assertEqual(xlsx.status_code, 200)
        workbook = load_workbook(BytesIO(xlsx.content))
        self.assertEqual(workbook.sheetnames, ['Summary', 'Programmes', 'Nominations'])
        programmes = list(workbook['Programmes'].values)
        self.assertEqual((programmes[1][1], programmes[1][-1]), ('IFMIS', 45000))
        nominations = list(workbook['Nominations'].values)
        self.assertEqual((nominations[1][1], nominations[1][10], nominations[1][11]), ('alice', 'Approved', 'Attended'))
        self.assertEqual(self.get(self.admin, f'/api/training-plans/{self.plan.pk}/export/?download=docx').status_code, 400)


def csv_file(text, name='staff.csv'):
    return SimpleUploadedFile(name, text.strip().encode('utf-8'), content_type='text/csv')


def xlsx_file(sheets, name='upload.xlsx'):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, rows in sheets.items():
        sheet = workbook.create_sheet(title)
        for row in rows:
            sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return SimpleUploadedFile(name, buffer.getvalue())


# An HRIS-style export: its own header names, a column LSMS doesn't know, and
# an officer whose supervisor is further down the file.
HRIS_EXPORT = """
PF No.,Surname,First Name,Official Email,Directorate,Designation,Job Group,Cadre,Reports To,Terms of Service,Date of Appointment,Sex,PWD,Station
1001,Otieno,Achieng,achieng@example.go.ke,Finance > Accounts,Accountant I,K,Accountants,1003,Permanent,01/07/2019,F,No,Kisumu
1002,Kamau,Brian,,Finance > Accounts,Accountant I,K,Accountants,1003,contract,2021-03-15,M,yes,HQ
1003,Wanjiru,Faith,faith@example.go.ke,Finance,Director Finance,R,Accountants,,P&P,,F,,HQ
"""


@FAST_HASHER
class StaffImportTests(MinistryTestCase):
    URL = '/api/imports/staff/'

    def upload(self, file, commit=False, user=None):
        self.client.force_authenticate(user=user or self.admin)
        return self.client.post(self.URL, {'file': file, 'commit': 'true' if commit else 'false'}, format='multipart')

    def test_check_reports_each_row_and_saves_nothing(self):
        response = self.upload(csv_file(HRIS_EXPORT))
        self.assertEqual(response.status_code, 200, response.content)
        data = response.json()
        self.assertEqual(data['counts']['created'], 3)
        self.assertEqual(data['extra_columns'], ['Station'])
        self.assertEqual(sorted(data['added']['grades']), ['K', 'R'])
        self.assertIn('Finance > Accounts', data['added']['departments'])
        self.assertFalse(User.objects.filter(staff_number='1001').exists())
        self.assertFalse(Grade.objects.filter(tenant=self.tenant).exists())
        self.assertFalse(ImportRun.objects.exists())

    def test_commit_creates_staff_in_their_posts_with_supervisors(self):
        data = self.upload(csv_file(HRIS_EXPORT), commit=True).json()
        self.assertEqual(data['counts']['created'], 3)
        achieng = User.objects.get(tenant=self.tenant, staff_number='1001')
        self.assertEqual((achieng.first_name, achieng.last_name, achieng.email), ('Achieng', 'Otieno', 'achieng@example.go.ke'))
        self.assertEqual((achieng.job_role.title, achieng.job_role.org_unit.name, achieng.job_role.org_unit.parent.name),
                         ('Accountant I', 'Accounts', 'Finance'))
        self.assertEqual((achieng.grade.code, achieng.job_family.name, achieng.supervisor.staff_number), ('K', 'Accountants', '1003'))
        self.assertEqual((achieng.employment_terms, achieng.date_of_appointment, achieng.gender, achieng.disability),
                         ('permanent', datetime.date(2019, 7, 1), 'female', False))
        self.assertEqual(achieng.extra, {'Station': 'Kisumu'})
        self.assertFalse(achieng.has_usable_password())
        # Existing Finance unit is reused, not duplicated, and the new post sits in the new Accounts division.
        self.assertEqual(OrgUnit.objects.filter(tenant=self.tenant, name='Finance').count(), 1)
        self.assertEqual(User.objects.get(staff_number='1002').job_role, achieng.job_role)
        self.assertEqual(ImportRun.objects.get().counts['created'], 3)

    def test_reimport_matches_on_staff_number_and_only_changes_what_changed(self):
        self.upload(csv_file(HRIS_EXPORT), commit=True)
        changed = HRIS_EXPORT.replace('1002,Kamau,Brian,,Finance > Accounts,Accountant I,K', '1002,Kamau,Brian,,Finance > Accounts,Accountant I,L')
        data = self.upload(csv_file(changed), commit=True).json()
        self.assertEqual((data['counts']['created'], data['counts']['updated'], data['counts']['unchanged']), (0, 1, 2))
        row = next(r for r in data['rows'] if r['key'] == '1002')
        self.assertEqual(row['changes'], ['grade'])
        self.assertEqual(User.objects.filter(tenant=self.tenant, staff_number__in=['1001', '1002', '1003']).count(), 3)

    def test_people_missing_from_the_file_are_listed_not_deleted(self):
        self.alice.staff_number = '0999'
        self.alice.save()
        data = self.upload(csv_file(HRIS_EXPORT), commit=True).json()
        self.assertEqual([p['staff_number'] for p in data['not_in_file']], ['0999'])
        self.assertTrue(User.objects.filter(pk=self.alice.pk, is_active=True).exists())

    def test_existing_account_is_matched_on_email_the_first_time(self):
        self.alice.email = 'achieng@example.go.ke'
        self.alice.save()
        self.upload(csv_file(HRIS_EXPORT), commit=True)
        self.alice.refresh_from_db()
        self.assertEqual((self.alice.staff_number, self.alice.first_name), ('1001', 'Achieng'))

    def test_bad_rows_are_reported_and_skipped(self):
        text = """
Staff number,First name,Last name,Department,Designation,Date of appointment,Gender
2001,Amina,Mohamed,HR,HR Officer,31/31/2020,female
,Kibet,Langat,HR,HR Officer,,
2003,Zawadi,Were,HR,HR Officer,,unknown
2004,Halima,Owino,HR,HR Officer,,
2004,Halima,Again,HR,HR Officer,,
"""
        data = self.upload(csv_file(text), commit=True).json()
        self.assertEqual((data['counts']['created'], data['counts']['error']), (1, 4))
        errors = {r['row']: r['errors'][0] for r in data['rows'] if r['action'] == 'error'}
        self.assertIn('is not a date', errors[2])
        self.assertIn('Staff number is empty', errors[3])
        self.assertIn('not a valid gender', errors[4])
        self.assertIn('already on row 5', errors[6])
        self.assertEqual(len(ImportRun.objects.get().errors), 4)

    def test_missing_required_column_rejects_the_file(self):
        response = self.upload(csv_file('Staff number,First name\n1,A'))
        self.assertEqual(response.status_code, 400)
        self.assertIn('Last name', response.json()['file'][0])

    def test_unsupported_file_is_rejected(self):
        response = self.upload(SimpleUploadedFile('staff.pdf', b'%PDF'))
        self.assertEqual(response.status_code, 400)

    def test_admin_cannot_change_their_own_access_by_import(self):
        self.admin.staff_number = '3001'
        self.admin.save()
        text = 'Staff number,First name,Last name,Department,Designation,LSMS access\n3001,Ad,Min,HR,HR Head,employee'
        data = self.upload(csv_file(text), commit=True).json()
        self.admin.refresh_from_db()
        self.assertEqual(self.admin.role, 'admin')
        self.assertTrue(data['rows'][0]['warnings'])

    def test_only_admins_import(self):
        for user in (self.manager, self.alice, self.superadmin):
            with self.subTest(user=user.username):
                self.assertEqual(self.upload(csv_file(HRIS_EXPORT), user=user).status_code, 403)

    def test_template_lists_the_columns_and_this_tenants_grades(self):
        Grade.objects.create(tenant=self.tenant, code='K', rank=1)
        response = self.get(self.admin, '/api/imports/staff/template/')
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(io.BytesIO(response.content))
        self.assertEqual(workbook.sheetnames, ['Instructions', 'Staff', 'Lists'])
        headers = [c.value for c in workbook['Staff'][1]]
        self.assertEqual(headers[:3], ['Staff number*', 'First name*', 'Last name*'])
        self.assertIn('K', [c.value for row in workbook['Lists'].iter_rows() for c in row])
        self.assertEqual(self.get(self.alice, '/api/imports/staff/template/').status_code, 403)

    def test_filled_template_imports(self):
        template = self.get(self.admin, '/api/imports/staff/template/').content
        workbook = load_workbook(io.BytesIO(template))
        workbook['Staff'].append(['4001', 'Mwende', 'Kilonzo', '', 'HR', 'HR Officer', 'J'])
        buffer = io.BytesIO()
        workbook.save(buffer)
        data = self.upload(SimpleUploadedFile('filled.xlsx', buffer.getvalue()), commit=True).json()
        self.assertEqual(data['counts']['created'], 1, data)


@FAST_HASHER
class CompetencyImportTests(MinistryTestCase):
    URL = '/api/imports/competencies/'
    COMPETENCIES = [
        ['Framework', 'Issuer', 'Code', 'Competency', 'Category', 'Difficulty', 'Importance', 'Frequency'],
        ['Demo Framework', 'Demo Commission', 'CC-01', 'Integrity', 'core', 2, 5, 5],
        ['Demo Framework', '', 'CC-02', 'Citizen focus', 'core', '', '', ''],
        ['Demo Framework', '', 'FC-01', 'IFMIS', 'technical', 4, 3, 4],
    ]
    REQUIREMENTS = [
        ['Designation', 'Competency', 'Required level'],
        ['Accountant', 'CC-01', 4],
        ['Accountant', 'Citizen focus', 'Intermediate'],
        ['Cleaner', 'CC-01', 3],
        ['HR Officer', 'CC-99', 7],
    ]

    def upload(self, sheets, commit=True):
        self.client.force_authenticate(user=self.admin)
        return self.client.post(self.URL, {'file': xlsx_file(sheets), 'commit': 'true' if commit else 'false'}, format='multipart')

    def test_framework_competencies_and_role_requirements(self):
        data = self.upload({'Competencies': self.COMPETENCIES, 'Role requirements': self.REQUIREMENTS}).json()
        self.assertEqual((data['counts']['created'], data['counts']['updated']), (2, 1))
        framework = CompetencyFramework.objects.get(tenant=self.tenant)
        self.assertEqual((framework.name, framework.issuer), ('Demo Framework', 'Demo Commission'))
        # The existing IFMIS skill joins the framework rather than being duplicated.
        self.ifmis.refresh_from_db()
        self.assertEqual((self.ifmis.framework, self.ifmis.code, self.ifmis.usage_frequency), (framework, 'FC-01', 4))
        self.assertEqual(Skill.objects.get(code='CC-02').difficulty, 3)
        requirements = data['requirements']
        self.assertEqual((requirements['counts']['created'], requirements['counts']['error']), (2, 2))
        errors = [r['errors'] for r in requirements['rows'] if r['action'] == 'error']
        self.assertIn('no post called "Cleaner"', errors[0][0])
        self.assertEqual(len(errors[1]), 2)
        integrity = JobRoleSkill.objects.get(job_role__title='Accountant', skill__code='CC-01')
        self.assertEqual(integrity.required_level, 4)
        self.assertEqual(JobRoleSkill.objects.get(job_role__title='Accountant', skill__code='CC-02').required_level, 3)

    def test_reimport_is_unchanged(self):
        sheets = {'Competencies': self.COMPETENCIES, 'Role requirements': self.REQUIREMENTS[:3]}
        self.upload(sheets)
        data = self.upload(sheets).json()
        self.assertEqual(data['counts']['unchanged'], 3)
        self.assertEqual(data['requirements']['counts']['unchanged'], 2)
        self.assertEqual(Skill.objects.filter(tenant=self.tenant).count(), 4)

    def test_csv_with_competencies_only(self):
        self.client.force_authenticate(user=self.admin)
        text = 'Framework,Code,Competency\nDemo,CC-01,Integrity'
        data = self.client.post(self.URL, {'file': csv_file(text, 'c.csv'), 'commit': 'true'}, format='multipart').json()
        self.assertEqual((data['counts']['created'], data['requirements']), (1, None))

    def test_template_has_both_sheets(self):
        response = self.get(self.admin, '/api/imports/competencies/template/')
        workbook = load_workbook(io.BytesIO(response.content))
        self.assertEqual(workbook.sheetnames, ['Instructions', 'Competencies', 'Role requirements', 'Lists'])


@FAST_HASHER
class StaffRecordTests(MinistryTestCase):
    def test_me_includes_the_tenants_wording(self):
        self.assertEqual(self.ok(self.alice, '/api/users/me/')['terms']['grade'], 'Grade')
        self.tenant.terminology = {'grade': 'Job Group'}
        self.tenant.save()
        terms = self.ok(self.alice, '/api/users/me/')['terms']
        self.assertEqual((terms['grade'], terms['staff_number']), ('Job Group', 'Staff number'))

    def test_new_organisation_can_start_from_a_pack(self):
        self.client.force_authenticate(user=self.superadmin)
        response = self.client.post('/api/tenants/', {'name': 'State Dept', 'pack': 'kenya_public_service'}, format='json')
        self.assertEqual(response.status_code, 201, response.content)
        tenant = Tenant.objects.get(name='State Dept')
        self.assertEqual(response.json()['terms']['grade'], 'Job Group')
        self.assertEqual(Grade.objects.filter(tenant=tenant).count(), 18)
        self.assertTrue(JobFamily.objects.filter(tenant=tenant, name='Accountants').exists())

    def test_staff_numbers_are_unique_within_a_tenant_only(self):
        self.client.force_authenticate(user=self.admin)
        self.assertEqual(self.client.patch(f'/api/users/{self.alice.pk}/', {'staff_number': 'P1'}, format='json').status_code, 200)
        response = self.client.patch(f'/api/users/{self.bob.pk}/', {'staff_number': 'p1'}, format='json')
        self.assertEqual(response.status_code, 400)
        self.outsider.staff_number = 'P2'
        self.outsider.save()
        self.assertEqual(self.client.patch(f'/api/users/{self.bob.pk}/', {'staff_number': 'P2'}, format='json').status_code, 200)

    def test_people_cannot_change_their_own_staff_record(self):
        self.client.force_authenticate(user=self.alice)
        response = self.client.patch(f'/api/users/{self.alice.pk}/', {'staff_number': 'X'}, format='json')
        self.assertEqual(response.status_code, 403)

    def test_supervisor_and_grade_must_be_in_the_tenant(self):
        other_grade = Grade.objects.create(tenant=self.other_tenant, code='Z')
        self.client.force_authenticate(user=self.admin)
        response = self.client.patch(f'/api/users/{self.alice.pk}/', {'grade': str(other_grade.pk), 'supervisor': str(self.outsider.pk)}, format='json')
        self.assertEqual(response.status_code, 400)
        self.assertEqual(set(response.json()), {'grade', 'supervisor'})

    def test_grades_are_tenant_scoped_and_admin_written(self):
        Grade.objects.create(tenant=self.other_tenant, code='Z')
        self.client.force_authenticate(user=self.admin)
        self.assertEqual(self.client.post('/api/grades/', {'code': 'K', 'name': 'Job Group K', 'rank': 9}, format='json').status_code, 201)
        self.assertEqual(self.client.post('/api/grades/', {'code': 'k'}, format='json').status_code, 400)
        self.assertEqual([g['code'] for g in self.ok(self.alice, '/api/grades/')], ['K'])
        self.client.force_authenticate(user=self.alice)
        self.assertEqual(self.client.post('/api/grades/', {'code': 'L'}, format='json').status_code, 403)

    def test_reports_show_the_staff_number(self):
        self.alice.staff_number = 'PF-123'
        self.alice.save()
        report = self.ok(self.admin, '/api/reports/tna/')
        alice = next(p for p in report['individuals'] if p['user']['username'] == 'alice')
        self.assertEqual(alice['user']['staff_number'], 'PF-123')
        bob = next(p for p in report['individuals'] if p['user']['username'] == 'bob')
        self.assertEqual(bob['user']['staff_number'], 'bob')


@FAST_HASHER
class DemoDataTests(TestCase):
    def test_seed_demo_loads_a_working_institution(self):
        call_command('seed_demo', stdout=io.StringIO())
        tenant = Tenant.objects.get(name='Demo State Department')
        self.assertEqual(tenant.terminology['grade'], 'Job Group')
        staff = User.objects.filter(tenant=tenant)
        self.assertGreater(staff.count(), 30)
        self.assertTrue(staff.filter(role='admin').exclude(password__startswith='!').exists())
        self.assertTrue(TrainingPlan.objects.get(tenant=tenant).programmes.exists())
        with self.assertRaises(Exception):
            call_command('seed_demo', stdout=io.StringIO())
        call_command('seed_demo', '--reset', stdout=io.StringIO())


def set_license(tenant, **fields):
    License.objects.filter(tenant=tenant).update(**fields)


def expired(days_ago):
    """Licence dates that ended ``days_ago`` days back."""
    today = datetime.date.today()
    return {'starts_on': today - timedelta(days=365), 'ends_on': today - timedelta(days=days_ago)}


@FAST_HASHER
class LicenseStateTests(TestCase):
    def test_new_organisation_starts_on_a_logged_pilot(self):
        tenant = Tenant.objects.create(name='New Agency')
        license = tenant.license
        self.assertEqual((license.plan, license.band, license.starts_on), ('pilot', 'small', datetime.date.today()))
        self.assertEqual(license.ends_on - license.starts_on, timedelta(days=licensing.PILOT_DAYS))
        self.assertEqual(LicenseEvent.objects.get(license=license).summary[:21], 'Pilot licence issued ')
        self.assertEqual(licensing.watermark(tenant), 'PILOT LICENCE - EVALUATION COPY')

    def test_state_runs_active_then_grace_then_read_only(self):
        license = License(starts_on=datetime.date(2026, 7, 1), ends_on=datetime.date(2027, 6, 30), grace_days=30)
        self.assertEqual(licensing.state(license, datetime.date(2027, 6, 30)), 'active')
        self.assertEqual(licensing.state(license, datetime.date(2027, 7, 30)), 'grace')
        self.assertEqual(licensing.state(license, datetime.date(2027, 7, 31)), 'read_only')
        license.suspended = True
        self.assertEqual(licensing.state(license, datetime.date(2027, 1, 1)), 'suspended')

    def test_financial_year_bounds(self):
        self.assertEqual(licensing.financial_year_bounds(datetime.date(2026, 10, 29)), (datetime.date(2026, 7, 1), datetime.date(2027, 6, 30)))
        self.assertEqual(licensing.financial_year_bounds(datetime.date(2027, 3, 1)), (datetime.date(2026, 7, 1), datetime.date(2027, 6, 30)))

    def test_every_api_view_checks_the_licence(self):
        from django.urls import get_resolver

        def views(patterns):
            for pattern in patterns:
                if hasattr(pattern, 'url_patterns'):
                    yield from views(pattern.url_patterns)
                else:
                    yield getattr(pattern.callback, 'cls', None)

        unchecked = {
            view.__name__ for view in views(get_resolver().url_patterns)
            if view is not None and view.__module__ == 'app.views'
            and licensing.HasWritableLicense not in view.permission_classes and view.__name__ != 'AcceptInvitationView'
        }
        self.assertEqual(unchecked, set())


class LicenseEnforcementTests(TenantTestCase):
    def post_skill(self, user):
        self.login(user)
        return self.client.post('/api/skills/', {'name': 'Budgeting', 'difficulty': 3, 'usage_frequency': 3, 'impact_if_missing': 3})

    def test_grace_period_keeps_full_access(self):
        set_license(self.tenant_a, **expired(days_ago=10))
        self.assertEqual(self.post_skill(self.admin_a).status_code, 201)

    def test_read_only_after_grace_blocks_changes_but_not_reading(self):
        set_license(self.tenant_a, **expired(days_ago=31))
        response = self.post_skill(self.admin_a)
        self.assertEqual(response.status_code, 403)
        self.assertIn('read-only', response.json()['detail'])
        self.assertEqual(self.client.get('/api/skills/').status_code, 200)
        self.assertEqual(self.client.get('/api/license/').json()['state'], 'read_only')
        # Other organisations are unaffected.
        self.assertEqual(self.post_skill(User.objects.create_user(username='admin_b', tenant=self.tenant_b, role='admin')).status_code, 201)

    def test_suspension_blocks_changes_with_the_reason(self):
        set_license(self.tenant_a, suspended=True, suspended_reason='invoice overdue')
        response = self.post_skill(self.admin_a)
        self.assertEqual(response.status_code, 403)
        self.assertIn('invoice overdue', response.json()['detail'])

    def test_super_admin_is_not_blocked(self):
        set_license(self.tenant_a, **expired(days_ago=60))
        root = User.objects.create_superuser(username='root', password=PASSWORD)
        self.login(root)
        response = self.client.post('/api/invitations/', {'email': 'x@acme.com', 'tenant': str(self.tenant_a.pk)})
        self.assertEqual(response.status_code, 201, response.content)

    def test_invitation_cannot_be_accepted_while_read_only(self):
        invitation = Invitation.objects.create(tenant=self.tenant_a, email='late@acme.com')
        set_license(self.tenant_a, **expired(days_ago=60))
        response = APIClient().post(f'/api/invitations/accept/{invitation.token}/', {'username': 'late', 'password': PASSWORD})
        self.assertEqual(response.status_code, 403)
        self.assertFalse(User.objects.filter(username='late').exists())

    def test_license_endpoint(self):
        set_license(self.tenant_a, lpo_number='LPO/123')
        self.login(self.user_a)
        data = self.client.get('/api/license/').json()
        self.assertEqual((data['plan'], data['state'], data['writable'], data['staff_count'], data['staff_limit'], data['lpo_number']),
                         ('pilot', 'active', True, 2, 300, 'LPO/123'))
        root = User.objects.create_superuser(username='root', password=PASSWORD)
        self.login(root)
        self.assertEqual(self.client.get('/api/license/').status_code, 400)
        self.assertEqual(self.client.get(f'/api/license/?tenant={self.tenant_b.pk}').json()['staff_count'], 1)


@mock.patch.dict(licensing.STAFF_LIMITS, {'small': 2})
class StaffLimitTests(TenantTestCase):
    """Acme has two active staff, which fills a band patched down to 2."""

    def test_invitation_refused_at_the_limit(self):
        self.login(self.admin_a)
        response = self.client.post('/api/invitations/', {'email': 'new@acme.com'})
        self.assertEqual(response.status_code, 403)
        self.assertIn('up to 2 active staff', response.json()['detail'])
        self.assertFalse(Invitation.objects.exists())

    def test_deactivating_someone_makes_room(self):
        self.user_a.is_active = False
        self.user_a.save()
        self.login(self.admin_a)
        self.assertEqual(self.client.post('/api/invitations/', {'email': 'new@acme.com'}).status_code, 201)

    def test_accepting_refused_at_the_limit(self):
        invitation = Invitation.objects.create(tenant=self.tenant_a, email='new@acme.com')
        response = APIClient().post(f'/api/invitations/accept/{invitation.token}/', {'username': 'newbie', 'password': PASSWORD})
        self.assertEqual(response.status_code, 403)
        self.assertFalse(User.objects.filter(username='newbie').exists())

    def test_adding_or_reactivating_a_user_refused_at_the_limit(self):
        self.login(self.admin_a)
        self.assertEqual(self.client.post('/api/users/', {'username': 'extra', 'password': PASSWORD}).status_code, 403)
        leaver = User.objects.create_user(username='leaver', tenant=self.tenant_a, is_active=False)
        self.assertEqual(self.client.patch(f'/api/users/{leaver.pk}/', {'is_active': True}).status_code, 403)
        leaver.refresh_from_db()
        self.assertFalse(leaver.is_active)


@FAST_HASHER
@mock.patch.dict(licensing.STAFF_LIMITS, {'small': 6})
class StaffImportLimitTests(MinistryTestCase):
    """The ministry has five staff; the HRIS export adds three, past a band patched down to 6."""

    def upload(self, commit):
        self.client.force_authenticate(user=self.admin)
        return self.client.post('/api/imports/staff/', {'file': csv_file(HRIS_EXPORT), 'commit': 'true' if commit else 'false'}, format='multipart')

    def test_check_warns_before_anything_is_saved(self):
        response = self.upload(commit=False)
        self.assertEqual(response.status_code, 403)
        self.assertIn('This change would bring it to 8.', response.json()['detail'])

    def test_commit_saves_nothing(self):
        self.assertEqual(self.upload(commit=True).status_code, 403)
        self.assertFalse(User.objects.filter(staff_number='1001').exists())
        self.assertFalse(ImportRun.objects.exists())

    def test_fits_once_the_band_is_larger(self):
        set_license(self.tenant, band='medium')
        self.assertEqual(self.upload(commit=True).status_code, 200)


class TNAReportWatermarkTests(MinistryTestCase):
    URL = '/api/reports/tna/'

    def test_pilot_report_is_watermarked(self):
        self.assertEqual(self.ok(self.admin, self.URL)['watermark'], 'PILOT LICENCE - EVALUATION COPY')
        sheet = load_workbook(io.BytesIO(self.get(self.admin, f'{self.URL}?download=xlsx').content))['Summary']
        self.assertEqual((sheet['A1'].value, sheet['A2'].value), ('PILOT LICENCE - EVALUATION COPY', 'Ministry'))
        pdf = self.get(self.admin, f'{self.URL}?download=pdf')
        self.assertEqual(pdf.status_code, 200)

    def test_paid_report_is_not(self):
        set_license(self.tenant, plan='standard')
        self.assertEqual(self.ok(self.admin, self.URL)['watermark'], '')
        sheet = load_workbook(io.BytesIO(self.get(self.admin, f'{self.URL}?download=xlsx').content))['Summary']
        self.assertEqual(sheet['A1'].value, 'Ministry')
