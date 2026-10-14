"""
The annual training plan: built from the TNA gaps, costed in KES, then tracked
through nomination, attendance and evaluation.

``generate_from_gaps`` turns the gap engine's training needs into one programme
per skill and nominates every officer with a gap in it, keeping the gap figures
as evidence. Nothing here recomputes a gap: the numbers come from GapEngine and
the priority bands from the TNA report.

``build_plan`` returns the plan as plain data (totals, cost by quarter,
department and priority, attendance and evaluation results, and the
programmes). ``render_pdf`` and ``render_xlsx`` turn that data into files.
"""

import io
import re
from collections import defaultdict
from decimal import Decimal

from django.db import transaction
from django.utils import timezone

from .models import QUARTERS, EmployeeSkillRating, TrainingNomination, TrainingProgramme
from .tna_report import LEVEL_NAMES, PRIORITY_NAMES, skill_need

QUARTER_NAMES = dict(QUARTERS)
# Higher-priority needs are trained earlier in the year; Q4 is left for catch-up and evaluation.
QUARTER_FOR_PRIORITY = {'high': 1, 'medium': 2, 'low': 3}
PRIORITY_ORDER = {'high': 0, 'medium': 1, 'low': 2}
MODE_NAMES = dict(TrainingProgramme.DELIVERY_CHOICES)
STATUS_NAMES = dict(TrainingNomination.STATUS_CHOICES)
ATTENDANCE_NAMES = dict(TrainingNomination.ATTENDANCE_CHOICES)


@transaction.atomic
def generate_from_gaps(plan, engine, nominated_by=None):
    """
    Adds a programme for every skill with gaps across the institution, and nominates
    each officer with a gap. Running it again only adds what is missing, so HR's
    edits, removed nominees and manual nominations stay as they are.
    """
    gaps_by_skill = defaultdict(list)
    for person in engine.people:
        for row in engine.skill_gaps(person):
            if row['gap']:
                gaps_by_skill[row['skill']['id']].append((person, row))

    programmes = {p.skill_id: p for p in plan.programmes.all() if p.skill_id}
    created = {'programmes': 0, 'nominations': 0}
    for skill in engine.summarise(engine.people)['skills']:
        entries = gaps_by_skill.get(skill['skill']['id'])
        if not entries:
            continue
        programme = programmes.get(skill['skill']['id'])
        if programme is None:
            priority = skill_need(skill)['priority']
            priority = priority if priority in QUARTER_FOR_PRIORITY else 'low'
            programme = TrainingProgramme.objects.create(
                tenant=plan.tenant, plan=plan, skill_id=skill['skill']['id'], title=skill['skill']['name'],
                priority=priority, quarter=QUARTER_FOR_PRIORITY[priority], duration_days=plan.default_duration_days,
                fee=plan.default_fee, allowance_per_day=plan.default_allowance_per_day,
            )
            created['programmes'] += 1
        nominated = set(programme.nominations.values_list('user_id', flat=True))
        for person, row in entries:
            if person.pk in nominated:
                continue
            TrainingNomination.objects.create(
                tenant=plan.tenant, programme=programme, user=person, source='gap', nominated_by=nominated_by,
                required_level=row['required_level'], current_level=row['current_level'], gap=row['gap'],
                priority_score=row['priority_score'],
            )
            created['nominations'] += 1
    return created


def record_manager_rating(user, skill, level):
    """
    Records a supervisor's rating of ``skill`` for ``user`` today, so the gap engine
    picks up what the training changed. Today's rating is updated if there is one.
    """
    today = timezone.localdate()
    ratings = EmployeeSkillRating.objects.filter(user=user, skill=skill)
    rating = ratings.filter(survey_date=today).first()
    if rating:
        rating.manager_rating = level
        rating.save(update_fields=['manager_rating'])
        return rating
    latest = ratings.order_by('-survey_date').first()
    return EmployeeSkillRating.objects.create(
        tenant=user.tenant, user=user, skill=skill, self_rating=latest.self_rating if latest else level,
        manager_rating=level, survey_date=today,
    )


# Figures

def money(value):
    return None if value is None else float(Decimal(value).quantize(Decimal('0.01')))


def _mean(values):
    values = [v for v in values if v is not None]
    return round(sum(values) / len(values), 2) if values else None


def _person(user):
    role = user.job_role
    return {
        'id': user.pk, 'username': user.username, 'name': user.get_full_name(),
        'job_role': {'id': role.pk, 'title': role.title} if role else None,
        'org_unit': {'id': role.org_unit.pk, 'name': role.org_unit.name} if role and role.org_unit else None,
    }


def active_nominations(programme):
    return [n for n in programme.nominations.all() if n.status in TrainingNomination.ACTIVE]


def programme_figures(programme):
    """Participants and cost of one programme. Declined and withdrawn nominees cost nothing."""
    participants = len(active_nominations(programme))
    return {
        'participants': participants,
        'cost_per_participant': money(programme.cost_per_participant),
        'total_cost': money(programme.fixed_cost + programme.cost_per_participant * participants),
    }


def evaluation_figures(nominations):
    evaluations = [n.evaluation for n in nominations if hasattr(n, 'evaluation')]
    participant = [e for e in evaluations if e.participant_submitted_at]
    tested = [e for e in evaluations if e.pre_test_score is not None and e.post_test_score is not None]
    supervisor = [(e, e.nomination) for e in evaluations if e.post_training_level is not None]
    return {
        'participant_responses': len(participant),
        'average_reaction': _mean([e.reaction_rating for e in participant]),
        'average_relevance': _mean([e.relevance_rating for e in participant]),
        'tested': len(tested),
        'average_learning_gain': _mean([e.post_test_score - e.pre_test_score for e in tested]),
        'supervisor_evaluations': len(supervisor),
        'gaps_closed': sum(1 for e, n in supervisor if n.required_level is not None and e.post_training_level >= n.required_level),
    }


def attendance_figures(nominations):
    approved = [n for n in nominations if n.status == 'approved']
    counts = {key: sum(1 for n in approved if n.attendance == key) for key in ATTENDANCE_NAMES}
    recorded = counts['attended'] + counts['partial'] + counts['absent']
    return {
        **counts,
        'attendance_rate': round(100 * (counts['attended'] + counts['partial']) / recorded, 1) if recorded else None,
    }


def programme_row(programme):
    nominations = list(programme.nominations.all())
    return {
        'id': programme.pk,
        'title': programme.title,
        'skill': {'id': programme.skill.pk, 'name': programme.skill.name} if programme.skill else None,
        'provider': programme.provider,
        'delivery_mode': programme.delivery_mode,
        'quarter': programme.quarter,
        'start_date': programme.start_date,
        'end_date': programme.end_date,
        'venue': programme.venue,
        'duration_days': programme.duration_days,
        'fee': money(programme.fee),
        'allowance_per_day': money(programme.allowance_per_day),
        'fixed_cost': money(programme.fixed_cost),
        'priority': programme.priority,
        **programme_figures(programme),
        'approved': sum(1 for n in nominations if n.status == 'approved'),
        'attended': sum(1 for n in nominations if n.status == 'approved' and n.attendance in ('attended', 'partial')),
    }


def plan_programmes(plan):
    programmes = list(
        plan.programmes.select_related('skill')
        .prefetch_related('nominations__user__job_role__org_unit', 'nominations__evaluation')
    )
    programmes.sort(key=lambda p: (p.quarter, PRIORITY_ORDER.get(p.priority, 3), p.title.lower()))
    return programmes


def build_plan(plan):
    """The plan with its totals, cost breakdowns, attendance and evaluation results, and programmes."""
    programmes = plan_programmes(plan)
    rows = [programme_row(p) for p in programmes]
    nominations = [n for p in programmes for n in p.nominations.all()]
    active = [n for n in nominations if n.status in TrainingNomination.ACTIVE]

    quarters = {q: {'quarter': q, 'label': QUARTER_NAMES[q], 'programmes': 0, 'participants': 0, 'cost': 0.0} for q in QUARTER_NAMES}
    priorities = {key: {'priority': key, 'programmes': 0, 'participants': 0, 'cost': 0.0} for key in PRIORITY_ORDER}
    departments = {}
    for programme, row in zip(programmes, rows):
        for group in (quarters[programme.quarter], priorities.setdefault(programme.priority, {'priority': programme.priority, 'programmes': 0, 'participants': 0, 'cost': 0.0})):
            group['programmes'] += 1
            group['participants'] += row['participants']
            group['cost'] += row['total_cost']
        # Each participant carries their own cost plus an equal share of the programme's fixed cost.
        share = row['total_cost'] / row['participants'] if row['participants'] else 0
        for nomination in active_nominations(programme):
            unit = _person(nomination.user)['org_unit']
            key = unit['id'] if unit else None
            department = departments.setdefault(key, {'org_unit': unit, 'participants': 0, 'cost': 0.0})
            department['participants'] += 1
            department['cost'] += share

    def rounded(groups):
        return [{**g, 'cost': round(g['cost'], 2)} for g in groups]

    total = round(sum(row['total_cost'] for row in rows), 2)
    budget = money(plan.budget)
    return {
        'id': plan.pk,
        'title': plan.title or f'Annual Training Plan {plan.period}',
        'institution': {'id': plan.tenant.pk, 'name': plan.tenant.name},
        'period': plan.period,
        'status': plan.status,
        'approved_at': plan.approved_at,
        'summary': {
            'programmes': len(rows),
            'participants': len(active),
            'officers': len({n.user_id for n in active}),
            'approved_participants': sum(1 for n in nominations if n.status == 'approved'),
            'total_cost': total,
            'budget': budget,
            'balance': None if budget is None else round(budget - total, 2),
        },
        'by_quarter': rounded(quarters.values()),
        'by_priority': rounded(p for p in priorities.values() if p['programmes']),
        'by_department': rounded(sorted(departments.values(), key=lambda d: (-d['cost'], d['org_unit']['name'] if d['org_unit'] else '~'))),
        'attendance': attendance_figures(nominations),
        'evaluation': evaluation_figures(nominations),
        'programmes': rows,
    }


def participants(plan):
    """Every nomination in the plan, programme by programme, for the export annex."""
    rows = []
    for programme in plan_programmes(plan):
        for n in sorted(programme.nominations.all(), key=lambda n: (-(n.priority_score or 0), n.user.username)):
            evaluation = getattr(n, 'evaluation', None)
            rows.append({'programme': programme, 'nomination': n, 'person': _person(n.user), 'evaluation': evaluation})
    return rows


# Shared table layouts for the PDF and the workbook

def filename(data, extension):
    slug = re.sub(r'[^A-Za-z0-9]+', '-', f"Training plan {data['institution']['name']} {data['period']}").strip('-')
    return f'{slug}.{extension}'


def kes(value):
    return '—' if value is None else f'KES {value:,.0f}'


def _dash(value):
    return '—' if value is None or value == '' else value


def _dates(row):
    if row['start_date'] and row['end_date']:
        return f"{row['start_date']:%d %b %Y} – {row['end_date']:%d %b %Y}"
    return f"{row['start_date']:%d %b %Y}" if row['start_date'] else '—'


def summary_rows(data):
    s, e, a = data['summary'], data['evaluation'], data['attendance']
    return [
        ('Training programmes', s['programmes']),
        ('Participants (nominations)', s['participants']),
        ('Officers to be trained', s['officers']),
        ('Nominations approved', s['approved_participants']),
        ('Total cost', kes(s['total_cost'])),
        ('Approved training budget', kes(s['budget'])),
        ('Balance', f"{kes(-s['balance'])} over budget" if s['balance'] is not None and s['balance'] < 0 else kes(s['balance'])),
        ('Attendance rate (approved nominees recorded)', '—' if a['attendance_rate'] is None else f"{a['attendance_rate']}%"),
        ('Average participant rating (1-5)', _dash(e['average_reaction'])),
        ('Average learning gain (test points)', _dash(e['average_learning_gain'])),
        ('Gaps closed after training', f"{e['gaps_closed']} of {e['supervisor_evaluations']} evaluated"),
    ]


def quarter_table(data):
    return ['Quarter', 'Programmes', 'Participants', 'Cost'], [[q['label'], q['programmes'], q['participants'], kes(q['cost'])] for q in data['by_quarter']]


def department_table(data):
    return ['Department', 'Participants', 'Cost'], [[d['org_unit']['name'] if d['org_unit'] else 'No department', d['participants'], kes(d['cost'])] for d in data['by_department']]


def programmes_table(data):
    header = ['#', 'Programme', 'Skill', 'Priority', 'Quarter', 'Mode', 'Provider / venue', 'Dates', 'Days', 'Participants', 'Cost per participant', 'Fixed cost', 'Total cost']
    rows = [[i, r['title'], r['skill']['name'] if r['skill'] else '—', PRIORITY_NAMES[r['priority']], QUARTER_NAMES[r['quarter']],
             MODE_NAMES[r['delivery_mode']], ' / '.join(v for v in (r['provider'], r['venue']) if v) or '—', _dates(r), r['duration_days'],
             r['participants'], kes(r['cost_per_participant']), kes(r['fixed_cost']), kes(r['total_cost'])]
            for i, r in enumerate(data['programmes'], 1)]
    return header, rows


def _level(value):
    return '—' if value is None else f'{value} {LEVEL_NAMES.get(value, "")}'.strip()


def participants_table(rows):
    header = ['Programme', 'Officer', 'P/F No.', 'Job role', 'Department', 'Current', 'Required', 'Status', 'Attendance', 'Rating (1-5)', 'Level after training']
    table = []
    for row in rows:
        n, p, e = row['nomination'], row['person'], row['evaluation']
        table.append([row['programme'].title, p['name'] or p['username'], p['username'], p['job_role']['title'] if p['job_role'] else '—',
                      p['org_unit']['name'] if p['org_unit'] else '—', _level(n.current_level), _level(n.required_level), STATUS_NAMES[n.status],
                      ATTENDANCE_NAMES[n.attendance], _dash(e.reaction_rating if e else None), _level(e.post_training_level if e else None)])
    return header, table


# PDF

def render_pdf(data, rows):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    body = ParagraphStyle('body', parent=styles['BodyText'], fontSize=9.5, leading=13)
    cell = ParagraphStyle('cell', parent=body, fontSize=8, leading=10)
    head_cell = ParagraphStyle('head', parent=cell, fontName='Helvetica-Bold', textColor=colors.white)
    h1 = ParagraphStyle('h1', parent=styles['Title'], fontSize=18, spaceAfter=4)
    h2 = ParagraphStyle('h2', parent=styles['Heading2'], fontSize=13, spaceBefore=10, spaceAfter=4)
    sub = ParagraphStyle('sub', parent=body, alignment=1, textColor=colors.HexColor('#444444'))
    navy = colors.HexColor('#1f3a5f')

    def esc(value):
        return str(value).replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')

    def table(header, rows, widths=None):
        if not rows:
            return Paragraph('<i>Nothing to report.</i>', body)
        t = Table([[Paragraph(esc(h), head_cell) for h in header]] + [[Paragraph(esc(v), cell) for v in row] for row in rows],
                  colWidths=widths, repeatRows=1, hAlign='LEFT')
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), navy),
            ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#b8c2cc')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f2f5f8')]),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('TOPPADDING', (0, 0), (-1, -1), 3),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ]))
        return t

    pagesize = landscape(A4)
    width = pagesize[0] - 30 * mm
    institution = data['institution']['name']

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont('Helvetica', 7.5)
        canvas.setFillColor(colors.HexColor('#666666'))
        canvas.drawString(15 * mm, 9 * mm, f"{institution} · Annual Training Plan {data['period']}")
        canvas.drawRightString(pagesize[0] - 15 * mm, 9 * mm, f'Page {doc.page}')
        canvas.restoreState()

    s = data['summary']
    over = s['balance'] is not None and s['balance'] < 0
    story = [
        Paragraph(esc(institution.upper()), sub),
        Spacer(1, 4),
        Paragraph(esc(data['title']), h1),
        Paragraph(esc(f"Financial year {data['period']} · {dict(draft='Draft', submitted='Submitted for approval', approved='Approved')[data['status']]}"), sub),
        Paragraph(esc(f"Generated {timezone.localdate():%d %B %Y} from LSMS"), sub),
        Spacer(1, 8),
        Paragraph('1. Summary', h2),
        Paragraph(esc(
            f"{s['programmes']} programmes will train {s['officers']} officers ({s['participants']} nominations) at a total cost of "
            f"{kes(s['total_cost'])}" + (f" against an approved budget of {kes(s['budget'])}" if s['budget'] is not None else '') +
            (', which exceeds the budget.' if over else '.') +
            ' Programmes and nominees come from the Training Needs Assessment for the same year.'), body),
        Spacer(1, 4),
        table(['Indicator', 'Value'], summary_rows(data), [width * 0.45, width * 0.25]),
        KeepTogether([Paragraph('2. Cost by quarter', h2), table(*quarter_table(data), [width * w for w in (0.3, 0.15, 0.15, 0.2)])]),
        KeepTogether([Paragraph('3. Cost by department', h2), table(*department_table(data), [width * w for w in (0.45, 0.15, 0.2)])]),
        KeepTogether([
            Paragraph('4. Training programmes', h2),
            Paragraph('Cost per participant is the course fee plus the daily allowance for each day of the course. '
                      'Declined and withdrawn nominees are not costed.', body),
            Spacer(1, 3),
        ]),
        table(*programmes_table(data), [width * w for w in (0.03, 0.12, 0.09, 0.055, 0.075, 0.07, 0.11, 0.1, 0.045, 0.08, 0.075, 0.065, 0.075)]),
        Spacer(1, 14),
        table(['', 'Name', 'Signature', 'Date'], [['Prepared by (Head of HRM&D)', '', '', ''], ['Approved by (Authorised Officer)', '', '', '']],
              [width * 0.3, width * 0.3, width * 0.2, width * 0.15]),
        PageBreak(),
        Paragraph('Annex: Nominated officers', h2),
        Paragraph('Officers nominated to each programme, with the gap from the TNA, their nomination status, attendance and evaluation.', body),
        Spacer(1, 3),
        table(*participants_table(rows), [width * w for w in (0.13, 0.12, 0.08, 0.1, 0.1, 0.08, 0.08, 0.07, 0.08, 0.07, 0.09)]),
    ]

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=pagesize, leftMargin=15 * mm, rightMargin=15 * mm, topMargin=14 * mm, bottomMargin=16 * mm,
                            title=f"{data['title']}", author=institution)
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buffer.getvalue()


# Excel

def render_xlsx(data, rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill('solid', fgColor='1F3A5F')
    wrap = Alignment(wrap_text=True, vertical='top')
    kes_format = '#,##0.00'

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Summary'
    sheet.append([data['institution']['name']])
    sheet.append([data['title']])
    sheet['A1'].font = Font(bold=True, size=12)
    sheet['A2'].font = Font(bold=True, size=14)
    sheet.append(['Financial year', data['period']])
    sheet.append(['Status', data['status'].title()])
    sheet.append([])

    def heading(ws, header):
        ws.append(header)
        for c in ws[ws.max_row]:
            c.font, c.fill, c.alignment = header_font, header_fill, wrap

    heading(sheet, ['Indicator', 'Value'])
    for label, value in summary_rows(data):
        sheet.append([label, value])
    sheet.append([])
    heading(sheet, ['Quarter', 'Programmes', 'Participants', 'Cost (KES)'])
    for q in data['by_quarter']:
        sheet.append([q['label'], q['programmes'], q['participants'], q['cost']])
        sheet.cell(sheet.max_row, 4).number_format = kes_format
    sheet.append([])
    heading(sheet, ['Department', 'Participants', 'Cost (KES)'])
    for d in data['by_department']:
        sheet.append([d['org_unit']['name'] if d['org_unit'] else 'No department', d['participants'], d['cost']])
        sheet.cell(sheet.max_row, 3).number_format = kes_format
    sheet.column_dimensions['A'].width = 46
    for column in 'BCD':
        sheet.column_dimensions[column].width = 18

    def add_sheet(title, header, table, money_columns=()):
        ws = workbook.create_sheet(title)
        heading(ws, header)
        for row in table:
            ws.append(row)
            for index in money_columns:
                ws.cell(ws.max_row, index).number_format = kes_format
        ws.freeze_panes = 'A2'
        if table:
            ws.auto_filter.ref = ws.dimensions
        for index, name in enumerate(header, 1):
            longest = max([len(str(name))] + [len(str(row[index - 1])) for row in table])
            ws.column_dimensions[get_column_letter(index)].width = min(max(longest + 2, 8), 50)

    # Costs stay numeric so the plan can be re-totalled in Excel.
    header = ['#', 'Programme', 'Skill', 'Priority', 'Quarter', 'Delivery', 'Provider', 'Venue', 'Start', 'End', 'Days',
              'Participants', 'Fee per participant', 'Allowance per day', 'Cost per participant', 'Fixed cost', 'Total cost (KES)']
    table = [[i, r['title'], r['skill']['name'] if r['skill'] else '', PRIORITY_NAMES[r['priority']], QUARTER_NAMES[r['quarter']],
              MODE_NAMES[r['delivery_mode']], r['provider'], r['venue'], r['start_date'], r['end_date'], r['duration_days'],
              r['participants'], r['fee'], r['allowance_per_day'], r['cost_per_participant'], r['fixed_cost'], r['total_cost']]
             for i, r in enumerate(data['programmes'], 1)]
    add_sheet('Programmes', header, table, money_columns=(13, 14, 15, 16, 17))

    header = ['Programme', 'Officer', 'P/F No.', 'Job role', 'Department', 'Source', 'Current level', 'Required level', 'Gap',
              'Priority score', 'Status', 'Attendance', 'Days attended', 'Reaction (1-5)', 'Relevance (1-5)', 'Pre-test', 'Post-test',
              'Level after training', 'Supervisor comments']
    table = []
    for row in rows:
        n, p, e = row['nomination'], row['person'], row['evaluation']
        table.append([row['programme'].title, p['name'] or p['username'], p['username'], p['job_role']['title'] if p['job_role'] else '',
                      p['org_unit']['name'] if p['org_unit'] else '', dict(TrainingNomination.SOURCE_CHOICES)[n.source], n.current_level,
                      n.required_level, n.gap, n.priority_score, STATUS_NAMES[n.status], ATTENDANCE_NAMES[n.attendance], n.days_attended,
                      e.reaction_rating if e else None, e.relevance_rating if e else None, e.pre_test_score if e else None,
                      e.post_test_score if e else None, e.post_training_level if e else None, e.supervisor_comments if e else ''])
    add_sheet('Nominations', header, table)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
