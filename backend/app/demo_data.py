"""
A made-up institution for demos: staff and a competency framework in the same
shape as the import templates, so the demo can be loaded through the importers
and the files handed to someone to try the upload themselves.

Every name, number and competency here is invented. The framework is a stand-in
for a real one (such as a public service commission's), which is loaded with
the competency template.
"""

import io
import random

from openpyxl import load_workbook

from . import imports

FRAMEWORK = 'Sample Public Service Competency Framework (demo data)'
ISSUER = 'Demo data: replace with the official framework'

# (code, name, category, description, difficulty, importance, frequency)
COMPETENCIES = [
    ('CC-01', 'Integrity', 'core', 'Acts honestly, declares conflicts of interest and upholds the values of public service.', 2, 5, 5),
    ('CC-02', 'Citizen focus', 'core', 'Understands what citizens need and delivers services courteously and on time.', 3, 5, 5),
    ('CC-03', 'Communication', 'core', 'Writes and speaks clearly to colleagues, stakeholders and the public.', 3, 4, 5),
    ('CC-04', 'Teamwork', 'core', 'Works with others across units to get shared results.', 2, 4, 5),
    ('CC-05', 'Results orientation', 'core', 'Plans work against performance targets and delivers them.', 3, 5, 4),
    ('CC-06', 'Digital literacy', 'core', 'Uses government digital systems and office software confidently and securely.', 3, 4, 5),
    ('CC-07', 'Innovation', 'core', 'Looks for better ways of delivering services and puts them into practice.', 4, 3, 3),
    ('LC-01', 'Strategic thinking', 'leadership', 'Links the unit\'s work to the institution\'s strategic plan and priorities.', 5, 5, 3),
    ('LC-02', 'Leading and developing people', 'leadership', 'Sets direction, coaches staff and manages their performance.', 4, 5, 4),
    ('LC-03', 'Decision making', 'leadership', 'Makes sound, timely and accountable decisions.', 4, 5, 4),
    ('FC-FIN-01', 'Public financial management', 'functional', 'Applies public finance law and regulations to budgeting and spending.', 4, 5, 4),
    ('FC-FIN-02', 'Financial reporting', 'technical', 'Prepares financial statements to the applicable accounting standards.', 4, 5, 4),
    ('FC-FIN-03', 'Budget preparation', 'technical', 'Prepares and tracks budgets against approved estimates.', 4, 4, 3),
    ('FC-HR-01', 'HR policies and procedures', 'functional', 'Applies the HR policies and procedures manual correctly.', 3, 5, 5),
    ('FC-HR-02', 'Performance management', 'technical', 'Runs staff appraisal, target setting and performance reviews.', 3, 4, 4),
    ('FC-HR-03', 'Training needs assessment', 'technical', 'Identifies skill gaps and plans training to close them.', 4, 4, 3),
    ('FC-ICT-01', 'Network administration', 'technical', 'Keeps the institution\'s networks running and available.', 4, 5, 5),
    ('FC-ICT-02', 'Cybersecurity', 'technical', 'Protects systems and data against threats and responds to incidents.', 5, 5, 4),
    ('FC-ICT-03', 'Systems support', 'technical', 'Supports users and maintains hardware and software.', 3, 4, 5),
    ('FC-SCM-01', 'Procurement law and regulations', 'functional', 'Applies procurement law and regulations to every purchase.', 4, 5, 5),
    ('FC-SCM-02', 'E-procurement', 'technical', 'Runs tenders and purchases through the e-procurement system.', 3, 4, 4),
    ('FC-REC-01', 'Records and archives management', 'technical', 'Files, tracks, retains and disposes of records correctly.', 3, 4, 5),
    ('FC-PLN-01', 'Policy analysis', 'functional', 'Analyses policy options and their effects.', 4, 4, 3),
    ('FC-PLN-02', 'Monitoring and evaluation', 'technical', 'Tracks programmes against indicators and reports on results.', 4, 4, 3),
    ('FC-PLN-03', 'Data analysis', 'technical', 'Cleans, analyses and presents data for decisions.', 4, 4, 4),
]

FUNCTIONAL = {
    'finance': ['FC-FIN-01', 'FC-FIN-02', 'FC-FIN-03'],
    'hr': ['FC-HR-01', 'FC-HR-02', 'FC-HR-03'],
    'ict': ['FC-ICT-01', 'FC-ICT-02', 'FC-ICT-03'],
    'scm': ['FC-SCM-01', 'FC-SCM-02'],
    'records': ['FC-REC-01'],
    'planning': ['FC-PLN-01', 'FC-PLN-02', 'FC-PLN-03'],
    'admin': [],
}
CORE = ['CC-01', 'CC-02', 'CC-03', 'CC-04', 'CC-05', 'CC-06', 'CC-07']
LEADERSHIP = ['LC-01', 'LC-02', 'LC-03']

# (department path, designation, job group, cadre, area, how many, is head of the unit)
POSTS = [
    ('Office of the Principal Secretary', 'Principal Secretary', 'T', 'Administrative Officers', 'admin', 1, True),
    ('Office of the Principal Secretary', 'Personal Secretary', 'J', 'Office Administrators', 'admin', 1, False),
    ('Finance and Accounts Directorate', 'Director, Finance', 'R', 'Accountants', 'finance', 1, True),
    ('Finance and Accounts Directorate > Accounts Division', 'Senior Accountant', 'L', 'Accountants', 'finance', 2, False),
    ('Finance and Accounts Directorate > Accounts Division', 'Accountant I', 'K', 'Accountants', 'finance', 3, False),
    ('Finance and Accounts Directorate > Budget Division', 'Economist I', 'K', 'Economists and Statisticians', 'finance', 2, False),
    ('Human Resource Management and Development Directorate', 'Director, HRM and D', 'R', 'Human Resource Management Officers', 'hr', 1, True),
    ('Human Resource Management and Development Directorate > HR Development Division', 'Senior HRM Officer', 'L', 'Human Resource Management Officers', 'hr', 2, False),
    ('Human Resource Management and Development Directorate > HR Development Division', 'HRM Officer I', 'K', 'Human Resource Management Officers', 'hr', 3, False),
    ('Human Resource Management and Development Directorate > Records Unit', 'Records Management Officer', 'J', 'Records Management Officers', 'records', 2, False),
    ('Human Resource Management and Development Directorate > Records Unit', 'Clerical Officer', 'F', 'Clerical Officers', 'records', 3, False),
    ('ICT Directorate', 'Director, ICT', 'R', 'ICT Officers', 'ict', 1, True),
    ('ICT Directorate', 'Senior ICT Officer', 'L', 'ICT Officers', 'ict', 2, False),
    ('ICT Directorate', 'ICT Officer I', 'K', 'ICT Officers', 'ict', 3, False),
    ('Supply Chain Management Division', 'Head, Supply Chain Management', 'P', 'Supply Chain Management Officers', 'scm', 1, True),
    ('Supply Chain Management Division', 'Supply Chain Management Officer', 'K', 'Supply Chain Management Officers', 'scm', 3, False),
    ('Planning Directorate', 'Director, Planning', 'R', 'Economists and Statisticians', 'planning', 1, True),
    ('Planning Directorate', 'Senior Economist', 'M', 'Economists and Statisticians', 'planning', 2, False),
    ('Planning Directorate', 'Statistician I', 'K', 'Economists and Statisticians', 'planning', 2, False),
]

FIRST_NAMES = [
    'Achieng', 'Wanjiru', 'Kiprono', 'Mutua', 'Nekesa', 'Omondi', 'Chebet', 'Kamau', 'Atieno', 'Mwende', 'Barasa', 'Njeri',
    'Otieno', 'Jepkosgei', 'Kariuki', 'Halima', 'Wafula', 'Nyambura', 'Kibet', 'Akinyi', 'Mumbua', 'Odhiambo', 'Wairimu',
    'Kipchumba', 'Zawadi', 'Musyoka', 'Adhiambo', 'Gitau', 'Amina', 'Cherono', 'Mwangi', 'Auma', 'Ndungu', 'Faith', 'Brian',
]
LAST_NAMES = [
    'Otieno', 'Kamau', 'Mutai', 'Wanyama', 'Njoroge', 'Ochieng', 'Kiptoo', 'Mwangi', 'Owino', 'Kilonzo', 'Wekesa', 'Maina',
    'Rotich', 'Onyango', 'Mohamed', 'Kimani', 'Chepkoech', 'Nyaga', 'Langat', 'Okoth', 'Muthoni', 'Simiyu', 'Koech', 'Were',
]
STATIONS = ['Headquarters', 'Headquarters', 'Headquarters', 'Mombasa regional office', 'Kisumu regional office', 'Nakuru regional office']

LEVEL_FOR_GRADE = {'T': 5, 'R': 4, 'P': 4, 'M': 4, 'L': 4}


def staff_rows(seed=7):
    """Rows for the staff template, plus the staff numbers of the demo's admin, manager and employee logins."""
    rng = random.Random(seed)
    rows, heads, logins = [], {}, {}
    number = 10001
    used_names = set()
    for path, title, grade, cadre, area, count, is_head in POSTS:
        top = path.split('>')[0].strip()
        for _ in range(count):
            while True:
                first, last = rng.choice(FIRST_NAMES), rng.choice(LAST_NAMES)
                if (first, last) not in used_names and first != last:
                    used_names.add((first, last))
                    break
            staff_number = f'DEMO{number}'
            number += 1
            if is_head:
                supervisor = heads.get('Office of the Principal Secretary', '') if top != 'Office of the Principal Secretary' else ''
                heads[top] = staff_number
            else:
                supervisor = heads.get(top, '')
            access = 'manager' if is_head else 'employee'
            if title == 'Director, HRM and D':
                access = 'admin'
                logins['admin'] = staff_number
            if title == 'Director, Finance':
                logins['manager'] = staff_number
            if title == 'Accountant I' and 'employee' not in logins:
                logins['employee'] = staff_number
            year = rng.randint(2004, 2023)
            rows.append({
                'staff_number': staff_number, 'first_name': first, 'last_name': last,
                'email': f'{first.lower()}.{last.lower()}@demo.example',
                'department': path, 'job_role': title, 'grade': grade, 'job_family': cadre,
                'supervisor': supervisor, 'role': access,
                'employment_terms': rng.choice(['permanent'] * 5 + ['contract']),
                'date_of_appointment': f'{year}-{rng.randint(1, 12):02d}-01',
                'gender': rng.choice(['female', 'male']),
                'disability': 'yes' if rng.random() < 0.05 else 'no',
                'Station': rng.choice(STATIONS),
            })
    return rows, logins


def competency_rows():
    return [
        {'framework': FRAMEWORK, 'issuer': ISSUER, 'version': 'Demo', 'code': code, 'name': name, 'type': category,
         'description': description, 'difficulty': d, 'impact_if_missing': i, 'usage_frequency': f}
        for code, name, category, description, d, i, f in COMPETENCIES
    ]


def requirement_rows():
    rows = []
    for path, title, grade, cadre, area, count, is_head in POSTS:
        senior = LEVEL_FOR_GRADE.get(grade, 3)
        codes = [(code, min(senior, 4)) for code in CORE]
        codes += [(code, senior) for code in FUNCTIONAL[area]]
        if is_head or grade in ('L', 'M'):
            codes += [(code, 4 if is_head else 3) for code in LEADERSHIP]
        rows += [{'job_role': title, 'department': path, 'competency': code, 'required_level': level} for code, level in codes]
    return rows


def _sheet(columns, rows):
    """Rows as an importer sheet: (headers, [(row number, values)]). Keys that aren't a column become extra columns."""
    fields = [c.field for c in columns]
    extra = [key for key in rows[0] if key not in fields] if rows else []
    headers = [c.header for c in columns] + extra
    return headers, [(n, [str(row.get(key, '')) for key in fields + extra]) for n, row in enumerate(rows, start=2)]


def staff_sheets():
    rows, _ = staff_rows()
    return {imports.STAFF_SHEET: _sheet(imports.STAFF_COLUMNS, rows)}


def competency_sheets():
    return {
        imports.COMPETENCY_SHEET: _sheet(imports.COMPETENCY_COLUMNS, competency_rows()),
        imports.REQUIREMENT_SHEET: _sheet(imports.REQUIREMENT_COLUMNS, requirement_rows()),
    }


def filled_template(template_bytes, sheets):
    """The import template with ``sheets``' rows filled in, as the bytes of an .xlsx file."""
    workbook = load_workbook(io.BytesIO(template_bytes))
    for name, (headers, rows) in sheets.items():
        sheet = workbook[name]
        for index, header in enumerate(headers, start=1):
            if sheet.cell(row=1, column=index).value is None:
                sheet.cell(row=1, column=index, value=header)
        for number, values in rows:
            for index, value in enumerate(values, start=1):
                sheet.cell(row=number, column=index, value=int(value) if value.isdigit() and len(value) < 3 else value)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
