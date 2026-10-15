"""
Spreadsheet imports: staff from an HR system export, and a competency framework
with the levels each job role requires.

Each import has a downloadable Excel template, but uploads don't have to use
it: column headers are matched loosely (case, spaces and punctuation are
ignored) against each column's known aliases, so an export from HRIS or
another HR system can usually be uploaded as it is. Columns the importer
doesn't know are kept on the person's record under ``extra``.

Every import runs twice: a check that changes nothing and reports what each row
would do, then the real run. Both go through the same code, inside a
transaction that the check rolls back, so the check can't disagree with the run.
"""

import csv
import datetime
import io
import re
from dataclasses import dataclass, field

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .models import CompetencyFramework, Grade, JobFamily, JobRole, JobRoleSkill, OrgUnit, Skill, PROFICIENCY_LEVELS
from .packs import terms_for

MAX_ROWS = 5000
MAX_FILE_BYTES = 5 * 1024 * 1024
PATH_SEPARATOR = '>'


class ImportFileError(Exception):
    """The file as a whole can't be read: wrong type, no header row, a required column missing."""


class EmptySheet(ImportFileError):
    pass


@dataclass
class Column:
    field: str
    header: str
    required: bool = False
    help: str = ''
    example: str = ''
    aliases: tuple = ()
    choices: tuple = ()  # Values offered as a drop-down in the template.


def normalise(text):
    return re.sub(r'[^a-z0-9]', '', str(text or '').lower())


STAFF_COLUMNS = [
    Column('staff_number', 'Staff number', True, 'The number HR knows the officer by. Used to match people on later imports.', 'KE-10234',
           ('pfno', 'pfnumber', 'personalnumber', 'personnelnumber', 'payrollnumber', 'employeenumber', 'employeeno', 'staffno', 'staffid', 'employeeid')),
    Column('first_name', 'First name', True, '', 'Achieng', ('firstname', 'givenname', 'othernames')),
    Column('last_name', 'Last name', True, '', 'Otieno', ('lastname', 'surname', 'familyname')),
    Column('email', 'Email', False, 'Work email, used to invite the officer to LSMS.', 'a.otieno@example.go.ke', ('emailaddress', 'officialemail', 'workemail')),
    Column('department', 'Department', True,
           f'Where the officer sits, from the top of the organisation down, separated by "{PATH_SEPARATOR}". Missing departments are created.',
           'Finance Directorate > Accounts Division', ('departmentpath', 'orgunit', 'division', 'section', 'unit', 'directorate')),
    Column('job_role', 'Designation', True, 'The post the officer holds. Missing posts are created in that department.', 'Senior Accountant',
           ('jobtitle', 'jobrole', 'position', 'post', 'title')),
    Column('grade', 'Grade', False, 'Grade code or name. New grades are created.', 'K', ('jobgroup', 'jg', 'gradecode', 'salarygrade', 'scale', 'band')),
    Column('job_family', 'Job family', False, 'Occupational group. New ones are created.', 'Accountants', ('cadre', 'schemeofservice', 'occupationalgroup')),
    Column('supervisor', 'Supervisor staff number', False, 'Staff number of the officer they report to.', 'KE-10001',
           ('supervisor', 'reportsto', 'supervisorpfno', 'supervisorno', 'supervisornumber', 'managerstaffnumber')),
    Column('role', 'LSMS access', False, 'employee (default), manager (can rate staff) or admin (HR).', 'employee',
           ('access', 'systemrole', 'lsmsrole'), ('employee', 'manager', 'admin')),
    Column('employment_terms', 'Terms of service', False, 'permanent, contract, temporary or intern.', 'permanent',
           ('terms', 'termsofemployment', 'employmentterms', 'termsofservice'), ('permanent', 'contract', 'temporary', 'intern')),
    Column('date_of_appointment', 'Date of appointment', False, 'As a date, or text such as 2019-07-01 or 01/07/2019.', '2019-07-01',
           ('appointmentdate', 'dateoffirstappointment', 'dateofemployment', 'datejoined')),
    Column('gender', 'Gender', False, 'female, male or other.', 'female', ('sex',), ('female', 'male', 'other')),
    Column('disability', 'Disability', False, 'yes or no.', 'no', ('pwd', 'personwithdisability', 'disabled'), ('yes', 'no')),
]

LEVEL_NAMES = ', '.join(f'{value} {label}' for value, label in PROFICIENCY_LEVELS)
CATEGORIES = ('core', 'functional', 'technical', 'leadership', 'soft')

COMPETENCY_COLUMNS = [
    Column('framework', 'Framework', True, 'Name of the framework the competency belongs to.', 'Public Service Competency Framework',
           ('frameworkname', 'competencyframework')),
    Column('issuer', 'Issuer', False, 'Who publishes the framework.', 'Public Service Commission', ('publisher', 'source')),
    Column('version', 'Version', False, '', '2024', ('edition',)),
    Column('code', 'Code', False, 'The competency\'s code in the framework. Used to match it on later imports.', 'CC-01',
           ('competencycode', 'ref', 'reference')),
    Column('name', 'Competency', True, '', 'Integrity', ('competencyname', 'name', 'skill', 'skillname', 'title')),
    Column('type', 'Category', False, 'core (default), functional, technical, leadership or soft.', 'core',
           ('competencytype', 'type', 'cluster', 'group'), CATEGORIES),
    Column('description', 'Description', False, 'What the competency means, or its behavioural indicators.',
           'Acts honestly and in line with the values of the public service.', ('definition', 'behaviouralindicators')),
    Column('difficulty', 'Difficulty', False, '1 (low) to 5 (high). Default 3.', '3', ('d',)),
    Column('impact_if_missing', 'Importance', False, '1 (low) to 5 (high). Default 3.', '4', ('i', 'impact', 'impactifmissing')),
    Column('usage_frequency', 'Frequency', False, '1 (low) to 5 (high). Default 3.', '5', ('f', 'frequencyofuse', 'usagefrequency')),
]

REQUIREMENT_COLUMNS = [
    Column('job_role', 'Designation', True, 'A post that already exists in LSMS (import staff first).', 'Senior Accountant', ('jobtitle', 'jobrole', 'position', 'post')),
    Column('department', 'Department', False, 'Only needed when posts in different departments share a title.', 'Finance Directorate > Accounts Division',
           ('departmentpath', 'orgunit', 'division')),
    Column('competency', 'Competency', True, 'The competency\'s code or name from the Competencies sheet.', 'CC-01', ('code', 'competencycode', 'competencyname', 'skill')),
    Column('required_level', 'Required level', True, f'1 to 5, or the level\'s name: {LEVEL_NAMES}.', '3', ('level', 'requiredproficiency', 'proficiency')),
    Column('difficulty', 'Difficulty', False, 'Only to differ from the competency\'s own score for this post.', '', ()),
    Column('importance', 'Importance', False, 'As above.', '', ('impact',)),
    Column('frequency', 'Frequency', False, 'As above.', '', ()),
]

STAFF_SHEET = 'Staff'
COMPETENCY_SHEET = 'Competencies'
REQUIREMENT_SHEET = 'Role requirements'


# Reading files

def _cell_text(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, (datetime.datetime, datetime.date)):
        return value
    return str(value).strip()


def _table(rows):
    """Turn a list of row tuples into (headers, [(row_number, values), ...]), skipping blank rows."""
    rows = iter(rows)
    for header_row in rows:
        headers = [str(h or '').strip().rstrip('*').strip() for h in header_row]
        if any(headers):
            break
    else:
        raise EmptySheet('The file is empty.')
    body = []
    for number, row in enumerate(rows, start=2):
        values = [_cell_text(v) for v in row]
        if any(v not in ('', None) for v in values):
            body.append((number, values))
        if len(body) > MAX_ROWS:
            raise ImportFileError(f'The file has more than {MAX_ROWS} rows. Split it into smaller files.')
    return headers, body


def read_file(upload):
    """
    Read an uploaded .xlsx or .csv into {sheet name: (headers, rows)}. A CSV
    becomes a single sheet named None; for a workbook the caller picks sheets by name.
    """
    if upload.size > MAX_FILE_BYTES:
        raise ImportFileError('The file is larger than 5 MB.')
    name = (upload.name or '').lower()
    data = upload.read()
    if name.endswith('.csv'):
        try:
            text = data.decode('utf-8-sig')
        except UnicodeDecodeError:
            text = data.decode('latin-1')
        return {None: _table(csv.reader(io.StringIO(text)))}
    if name.endswith('.xlsx'):
        try:
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception:
            raise ImportFileError('The file could not be opened as an Excel workbook.')
        sheets = {}
        for sheet in workbook.worksheets:
            if sheet.title in ('Instructions', 'Lists'):
                continue
            try:
                sheets[sheet.title] = _table(sheet.iter_rows(values_only=True))
            except EmptySheet:
                continue
        if not sheets:
            raise EmptySheet('The workbook has no data.')
        return sheets
    raise ImportFileError('Upload an Excel (.xlsx) or CSV (.csv) file.')


def pick_sheet(sheets, name, required=True):
    """The sheet called ``name`` (any case), else the only or first sheet when it's the one required."""
    for title, table in sheets.items():
        if title is not None and normalise(title) == normalise(name):
            return table
    if required and sheets:
        return next(iter(sheets.values()))
    return None


def map_headers(headers, columns):
    """Return ({field: column index}, [unknown headers with their index]); raise if a required column is missing."""
    lookup = {}
    for column in columns:
        for key in (column.field, column.header, *column.aliases):
            lookup.setdefault(normalise(key), column.field)
    mapping, unknown = {}, []
    for index, header in enumerate(headers):
        field_name = lookup.get(normalise(header))
        if field_name and field_name not in mapping:
            mapping[field_name] = index
        elif header:
            unknown.append((header, index))
    missing = [c.header for c in columns if c.required and c.field not in mapping]
    if missing:
        raise ImportFileError(f'The file has no {", ".join(missing)} column. Download the template to see the columns it needs.')
    return mapping, unknown


def _value(values, mapping, field_name):
    index = mapping.get(field_name)
    if index is None or index >= len(values):
        return ''
    return values[index]


# Parsing values

TERMS_VALUES = {
    'permanent': 'permanent', 'permanentandpensionable': 'permanent', 'pp': 'permanent', 'pandp': 'permanent',
    'contract': 'contract', 'contractual': 'contract',
    'temporary': 'temporary', 'temp': 'temporary', 'casual': 'temporary',
    'intern': 'intern', 'internship': 'intern', 'attachment': 'intern',
}
GENDER_VALUES = {'female': 'female', 'f': 'female', 'male': 'male', 'm': 'male', 'other': 'other'}
YES_NO = {'yes': True, 'y': True, 'true': True, '1': True, 'no': False, 'n': False, 'false': False, '0': False}
ROLES = {'employee': 'employee', 'staff': 'employee', 'officer': 'employee', 'manager': 'manager', 'supervisor': 'manager',
         'admin': 'admin', 'hr': 'admin', 'administrator': 'admin'}
DATE_FORMATS = ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y', '%d.%m.%Y', '%d %b %Y', '%d %B %Y')


def parse_date(value):
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    for fmt in DATE_FORMATS:
        try:
            return datetime.datetime.strptime(value, fmt).date()
        except ValueError:
            continue
    raise ValueError(f'"{value}" is not a date. Use 2019-07-01 or 01/07/2019.')


def parse_choice(value, options, label):
    key = normalise(value)
    if key not in options:
        raise ValueError(f'"{value}" is not a valid {label}.')
    return options[key]


def parse_score(value, label, allow_blank=True):
    if value == '' and allow_blank:
        return None
    try:
        score = int(str(value))
    except ValueError:
        raise ValueError(f'{label} must be a whole number from 1 to 5, not "{value}".')
    if not 1 <= score <= 5:
        raise ValueError(f'{label} must be from 1 to 5, not {score}.')
    return score


def parse_level(value):
    names = {normalise(label): level for level, label in PROFICIENCY_LEVELS}
    if normalise(value) in names:
        return names[normalise(value)]
    return parse_score(value, 'Required level', allow_blank=False)


def split_path(text):
    return [part.strip() for part in str(text).split(PATH_SEPARATOR) if part.strip()]


# Running an import

@dataclass
class RowResult:
    row: int
    key: str = ''
    label: str = ''
    action: str = 'unchanged'  # created, updated, unchanged or error
    changes: list = field(default_factory=list)
    errors: list = field(default_factory=list)
    warnings: list = field(default_factory=list)

    def as_dict(self):
        return {'row': self.row, 'key': self.key, 'label': self.label, 'action': self.action,
                'changes': self.changes, 'errors': self.errors, 'warnings': self.warnings}


class Importer:
    """Shared bookkeeping: the row results, the records an import adds along the way, and the check-or-commit transaction."""
    kind = ''

    def __init__(self, tenant, by=None):
        self.tenant = tenant
        self.by = by
        self.results = []
        self.added = {}  # What the import created besides its main records, by type, for the check to list.

    def note_added(self, kind, label):
        self.added.setdefault(kind, [])
        if label not in self.added[kind]:
            self.added[kind].append(label)

    def run(self, sheets, commit):
        with transaction.atomic():
            extra = self.process(sheets)
            if not commit:
                transaction.set_rollback(True)
        counts = {action: sum(1 for r in self.results if r.action == action) for action in ('created', 'updated', 'unchanged', 'error')}
        counts.update({f'{kind}_added': len(labels) for kind, labels in self.added.items()})
        return {
            'kind': self.kind,
            'committed': commit,
            'counts': counts,
            'rows': [r.as_dict() for r in self.results],
            'added': self.added,
            **extra,
        }

    def process(self, sheets):
        raise NotImplementedError


class StaffImporter(Importer):
    """
    Staff matched on staff number (or, the first time, on email), placed in
    their department and post. Departments, posts, grades and job families the
    file mentions but LSMS doesn't have yet are created. Nobody is deleted:
    people missing from the file are listed for HR to deactivate.
    """
    kind = 'staff'

    def __init__(self, tenant, by=None):
        super().__init__(tenant, by)
        self.User = get_user_model()
        self.units, self.units_by_name = {}, {}
        for unit in OrgUnit.objects.filter(tenant=tenant):
            self.units.setdefault((unit.parent_id, unit.name.lower()), unit)
            self.units_by_name.setdefault(unit.name.lower(), []).append(unit)
        self.roles = {(r.org_unit_id, r.title.lower()): r for r in JobRole.objects.filter(tenant=tenant)}
        self.grades = {}
        for grade in Grade.objects.filter(tenant=tenant):
            self._index_grade(grade)
        self.families = {}
        for family in JobFamily.objects.filter(tenant=tenant):
            self.families[normalise(family.name)] = family
            if family.code:
                self.families.setdefault(normalise(family.code), family)

    def _index_grade(self, grade):
        for key in (grade.code, grade.name, *(grade.aliases or [])):
            if key:
                self.grades.setdefault(normalise(key), grade)

    def unit_for(self, path):
        """
        The org unit at ``path``, creating what's missing. The path can start
        anywhere in the tree: if its first part isn't a top-level unit but names
        exactly one unit elsewhere (an export that gives only "Finance"), it
        starts from there.
        """
        parent, names = None, split_path(path)
        for depth, name in enumerate(names, start=1):
            key = (parent.pk if parent else None, name.lower())
            unit = self.units.get(key)
            if unit is None and depth == 1 and len(self.units_by_name.get(name.lower(), [])) == 1:
                unit = self.units_by_name[name.lower()][0]
            if unit is None:
                unit = self.units[key] = OrgUnit.objects.create(tenant=self.tenant, name=name, parent=parent)
                self.note_added('departments', f' {PATH_SEPARATOR} '.join(names[:depth]))
            parent = unit
        return parent

    def grade_for(self, text):
        grade = self.grades.get(normalise(text))
        if grade is None:
            rank = (Grade.objects.filter(tenant=self.tenant).order_by('-rank').values_list('rank', flat=True).first() or 0) + 1
            grade = Grade.objects.create(tenant=self.tenant, code=str(text)[:20], rank=rank)
            self._index_grade(grade)
            self.note_added('grades', grade.code)
        return grade

    def family_for(self, text):
        family = self.families.get(normalise(text))
        if family is None:
            family = JobFamily.objects.create(tenant=self.tenant, name=text)
            self.families[normalise(text)] = family
            self.note_added('job_families', text)
        return family

    def role_for(self, title, unit, grade, family):
        key = (unit.pk if unit else None, title.lower())
        role = self.roles.get(key)
        if role is None:
            role = JobRole.objects.create(tenant=self.tenant, org_unit=unit, title=title, grade=grade, job_family=family)
            self.roles[key] = role
            self.note_added('job_roles', f'{title} ({unit.name})' if unit else title)
        return role

    def unique_username(self, staff_number):
        base = re.sub(r'[^\w.@+-]', '', staff_number) or 'officer'
        username, n = base, 1
        while self.User.objects.filter(username__iexact=username).exists():
            n += 1
            username = f'{base}.{n}'
        return username

    def parse_row(self, values, mapping, result):
        """Read and check one row's values; errors go on ``result``."""
        get = lambda name: _value(values, mapping, name)  # noqa: E731
        data = {}
        for name in ('staff_number', 'first_name', 'last_name', 'department', 'job_role'):
            text = str(get(name))
            if not text:
                result.errors.append(f'{self.header(name)} is empty.')
            data[name] = text
        email = str(get('email'))
        if email:
            try:
                validate_email(email)
                data['email'] = email
            except ValidationError:
                result.errors.append(f'"{email}" is not an email address.')
        parsers = {
            'role': lambda v: parse_choice(v, ROLES, 'LSMS access (use employee, manager or admin)'),
            'employment_terms': lambda v: parse_choice(v, TERMS_VALUES, 'term of service'),
            'date_of_appointment': parse_date,
            'gender': lambda v: parse_choice(v, GENDER_VALUES, 'gender'),
            'disability': lambda v: parse_choice(v, YES_NO, 'disability answer (use yes or no)'),
        }
        for name, parse in parsers.items():
            value = get(name)
            if value not in ('', None):
                try:
                    data[name] = parse(value)
                except ValueError as exc:
                    result.errors.append(str(exc))
        for name in ('grade', 'job_family', 'supervisor'):
            if get(name):
                data[name] = str(get(name))
        return data

    def header(self, name):
        return next(c.header for c in STAFF_COLUMNS if c.field == name)

    def process(self, sheets):
        headers, rows = pick_sheet(sheets, STAFF_SHEET)
        mapping, unknown = map_headers(headers, STAFF_COLUMNS)
        terms = terms_for(self.tenant)
        staff = self.User.objects.filter(tenant=self.tenant)
        by_number = {u.staff_number.lower(): u for u in staff if u.staff_number}
        by_email = {u.email.lower(): u for u in staff if u.email and not u.staff_number}
        seen, pending_supervisors = {}, []

        for number, values in rows:
            result = RowResult(row=number)
            self.results.append(result)
            data = self.parse_row(values, mapping, result)
            result.key = data['staff_number']
            result.label = f"{data['first_name']} {data['last_name']}".strip()
            key = data['staff_number'].lower()
            if key and key in seen:
                result.errors.append(f'{terms["staff_number"]} {data["staff_number"]} is already on row {seen[key]}.')
            if result.errors:
                result.action = 'error'
                continue
            seen[key] = number

            unit = self.unit_for(data['department'])
            grade = self.grade_for(data['grade']) if data.get('grade') else None
            family = self.family_for(data['job_family']) if data.get('job_family') else None
            role = self.role_for(data['job_role'], unit, grade, family)

            person = by_number.get(key) or (by_email.get(data['email'].lower()) if data.get('email') else None)
            fields = {
                'staff_number': data['staff_number'],
                'first_name': data['first_name'],
                'last_name': data['last_name'],
                'job_role': role,
                **{name: data[name] for name in ('email', 'employment_terms', 'date_of_appointment', 'gender', 'disability') if name in data},
            }
            if grade:
                fields['grade'] = grade
            if family:
                fields['job_family'] = family
            if 'role' in data:
                if person is not None and person.pk == getattr(self.by, 'pk', None) and data['role'] != person.role:
                    result.warnings.append("Your own LSMS access wasn't changed.")
                else:
                    fields['role'] = data['role']
            extra = {header: str(values[index]) for header, index in unknown if index < len(values) and values[index] not in ('', None)}

            if person is None:
                person = self.User(tenant=self.tenant, username=self.unique_username(data['staff_number']), **fields)
                person.extra = extra
                person.set_unusable_password()
                person.save()
                result.action = 'created'
            else:
                for name, value in fields.items():
                    if getattr(person, name) != value:
                        result.changes.append(name)
                        setattr(person, name, value)
                merged = {**(person.extra or {}), **extra}
                if merged != (person.extra or {}):
                    result.changes.append('extra')
                    person.extra = merged
                if result.changes:
                    person.save()
                    result.action = 'updated'
            by_number[key] = person
            if data.get('supervisor'):
                pending_supervisors.append((result, person, data['supervisor']))

        # Supervisors last, so an officer can report to someone further down the file.
        for result, person, supervisor_number in pending_supervisors:
            supervisor = by_number.get(supervisor_number.lower())
            if supervisor is None:
                result.warnings.append(f'Supervisor {supervisor_number} is not in the file or in LSMS, so it was left blank.')
            elif supervisor.pk == person.pk:
                result.warnings.append('An officer can\'t supervise themselves, so the supervisor was left blank.')
            elif person.supervisor_id != supervisor.pk:
                person.supervisor = supervisor
                person.save(update_fields=['supervisor'])
                if result.action == 'unchanged':
                    result.action = 'updated'
                if result.action == 'updated':
                    result.changes.append('supervisor')

        in_file = {r.key.lower() for r in self.results if r.key}
        missing = (staff.filter(is_active=True).exclude(staff_number='')
                   .order_by('staff_number').values('id', 'staff_number', 'first_name', 'last_name'))
        return {
            'not_in_file': [
                {'id': p['id'], 'staff_number': p['staff_number'], 'name': f"{p['first_name']} {p['last_name']}".strip()}
                for p in missing if p['staff_number'].lower() not in in_file
            ],
            'extra_columns': [header for header, _ in unknown],
        }


class CompetencyImporter(Importer):
    """
    A competency framework's competencies, matched on code within their
    framework (or on name where there's no code), and optionally the level each
    job role requires, which becomes the role's skill profile.
    """
    kind = 'competencies'

    def process(self, sheets):
        headers, rows = pick_sheet(sheets, COMPETENCY_SHEET)
        mapping, _ = map_headers(headers, COMPETENCY_COLUMNS)
        frameworks = {normalise(f.name): f for f in CompetencyFramework.objects.filter(tenant=self.tenant)}
        skills = list(Skill.objects.filter(tenant=self.tenant))
        by_code = {(s.framework_id, s.code.lower()): s for s in skills if s.code}
        by_name = {}
        for skill in skills:
            by_name.setdefault(skill.name.lower(), skill)
        seen = {}

        for number, values in rows:
            get = lambda name: _value(values, mapping, name)  # noqa: E731
            result = RowResult(row=number, key=str(get('code')), label=str(get('name')))
            self.results.append(result)
            framework_name = str(get('framework'))
            for name, label in (('framework', 'Framework'), ('name', 'Competency')):
                if not get(name):
                    result.errors.append(f'{label} is empty.')
            scores = {}
            for name, label in (('difficulty', 'Difficulty'), ('impact_if_missing', 'Importance'), ('usage_frequency', 'Frequency')):
                try:
                    scores[name] = parse_score(get(name), label) or 3
                except ValueError as exc:
                    result.errors.append(str(exc))
            category = 'core'
            if get('type'):
                try:
                    category = parse_choice(get('type'), {c: c for c in CATEGORIES}, 'category (use core, functional, technical, leadership or soft)')
                except ValueError as exc:
                    result.errors.append(str(exc))
            ident = (normalise(framework_name), (result.key or result.label).lower())
            if ident in seen:
                result.errors.append(f'This competency is already on row {seen[ident]}.')
            if result.errors:
                result.action = 'error'
                continue
            seen[ident] = number

            framework = frameworks.get(normalise(framework_name))
            if framework is None:
                framework = CompetencyFramework.objects.create(
                    tenant=self.tenant, name=framework_name, issuer=str(get('issuer')), version=str(get('version')))
                frameworks[normalise(framework_name)] = framework
                self.note_added('frameworks', framework_name)

            skill = by_code.get((framework.pk, result.key.lower())) if result.key else None
            if skill is None:
                candidate = by_name.get(result.label.lower())
                if candidate is not None and candidate.framework_id in (None, framework.pk) and not (result.key and candidate.code):
                    skill = candidate
            fields = {'name': result.label, 'framework': framework, 'code': result.key, 'type': category, **scores}
            if get('description'):
                fields['description'] = str(get('description'))
            if skill is None:
                skill = Skill.objects.create(tenant=self.tenant, **fields)
                result.action = 'created'
            else:
                for name, value in fields.items():
                    if getattr(skill, name) != value:
                        result.changes.append(name)
                        setattr(skill, name, value)
                if result.changes:
                    skill.save()
                    result.action = 'updated'
            if skill.code:
                by_code[(framework.pk, skill.code.lower())] = skill
            by_name[skill.name.lower()] = skill

        requirements = pick_sheet(sheets, REQUIREMENT_SHEET, required=False)
        return {'requirements': self.process_requirements(requirements, by_code, by_name) if requirements else None}

    def process_requirements(self, table, by_code, by_name):
        headers, rows = table
        mapping, _ = map_headers(headers, REQUIREMENT_COLUMNS)
        roles = list(JobRole.objects.filter(tenant=self.tenant).select_related('org_unit'))
        codes = {}
        for (framework_id, code), skill in by_code.items():
            codes.setdefault(code, skill)
        results = []
        for number, values in rows:
            get = lambda name: _value(values, mapping, name)  # noqa: E731
            title, competency = str(get('job_role')), str(get('competency'))
            result = RowResult(row=number, key=competency, label=title)
            results.append(result)
            matches = [r for r in roles if r.title.lower() == title.lower()]
            if get('department'):
                unit_name = split_path(get('department'))[-1].lower() if split_path(get('department')) else ''
                matches = [r for r in matches if r.org_unit and r.org_unit.name.lower() == unit_name]
            if not title:
                result.errors.append('Designation is empty.')
            elif not matches:
                result.errors.append(f'There is no post called "{title}"{" in that department" if get("department") else ""}. Import staff first, or add the post.')
            elif len(matches) > 1:
                result.errors.append(f'More than one department has a "{title}" post. Add the Department column to say which.')
            skill = codes.get(competency.lower()) or by_name.get(competency.lower())
            if not competency:
                result.errors.append('Competency is empty.')
            elif skill is None:
                result.errors.append(f'There is no competency "{competency}".')
            fields = {}
            try:
                fields['required_level'] = parse_level(get('required_level'))
            except ValueError as exc:
                result.errors.append(str(exc))
            for name, label in (('difficulty', 'Difficulty'), ('importance', 'Importance'), ('frequency', 'Frequency')):
                try:
                    score = parse_score(get(name), label)
                except ValueError as exc:
                    result.errors.append(str(exc))
                else:
                    if score is not None:
                        fields[name] = score
            if result.errors:
                result.action = 'error'
                continue
            requirement = JobRoleSkill.objects.filter(job_role=matches[0], skill=skill).first()
            if requirement is None:
                JobRoleSkill.objects.create(tenant=self.tenant, job_role=matches[0], skill=skill, **fields)
                result.action = 'created'
            else:
                result.changes = [name for name, value in fields.items() if getattr(requirement, name) != value]
                if result.changes:
                    for name in result.changes:
                        setattr(requirement, name, fields[name])
                    requirement.save()
                    result.action = 'updated'
        counts = {action: sum(1 for r in results if r.action == action) for action in ('created', 'updated', 'unchanged', 'error')}
        return {'counts': counts, 'rows': [r.as_dict() for r in results]}


IMPORTERS = {'staff': StaffImporter, 'competencies': CompetencyImporter}


# Templates

HEADER_FILL = PatternFill('solid', fgColor='1F4E79')
REQUIRED_FILL = PatternFill('solid', fgColor='9C2A2A')
WHITE_BOLD = Font(bold=True, color='FFFFFF')


def _data_sheet(workbook, title, columns, lists):
    sheet = workbook.create_sheet(title)
    for index, column in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=index, value=f'{column.header}*' if column.required else column.header)
        cell.font = WHITE_BOLD
        cell.fill = REQUIRED_FILL if column.required else HEADER_FILL
        letter = cell.column_letter
        sheet.column_dimensions[letter].width = max(16, len(column.header) + 6)
        source = lists.get(column.field)
        if source:
            validation = DataValidation(type='list', formula1=source, allow_blank=True, showErrorMessage=False)
            validation.add(f'{letter}2:{letter}{MAX_ROWS + 1}')
            sheet.add_data_validation(validation)
    sheet.freeze_panes = 'A2'
    return sheet


def _instructions(sheet, title, intro, sections):
    sheet.column_dimensions['A'].width = 26
    sheet.column_dimensions['B'].width = 11
    sheet.column_dimensions['C'].width = 70
    sheet.column_dimensions['D'].width = 34
    sheet.append([title])
    sheet['A1'].font = Font(bold=True, size=14)
    for line in intro:
        sheet.append([line])
    for heading, columns in sections:
        sheet.append([])
        sheet.append([f'Sheet "{heading}"'])
        sheet.cell(row=sheet.max_row, column=1).font = Font(bold=True, size=12)
        sheet.append(['Column', 'Required', 'What to enter', 'Example'])
        for cell in sheet[sheet.max_row]:
            cell.font = WHITE_BOLD
            cell.fill = HEADER_FILL
        for column in columns:
            sheet.append([column.header, 'Yes' if column.required else '', column.help, column.example])
            sheet.cell(row=sheet.max_row, column=3).alignment = Alignment(wrap_text=True, vertical='top')


def _lists_sheet(workbook, lists):
    """A hidden sheet of drop-down values; returns {field: formula} for data validation."""
    sheet = workbook.create_sheet('Lists')
    formulas = {}
    for index, (name, values) in enumerate(lists.items(), start=1):
        if not values:
            continue
        for row, value in enumerate(values, start=1):
            cell = sheet.cell(row=row, column=index, value=value)
            letter = cell.column_letter
        formulas[name] = f"Lists!${letter}$1:${letter}${len(values)}"
    sheet.sheet_state = 'hidden'
    return formulas


def _save(workbook):
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def staff_template(tenant):
    terms = terms_for(tenant)
    workbook = Workbook()
    instructions = workbook.active
    instructions.title = 'Instructions'
    lists = _lists_sheet(workbook, {
        **{c.field: list(c.choices) for c in STAFF_COLUMNS if c.choices},
        'grade': list(Grade.objects.filter(tenant=tenant).values_list('code', flat=True)),
        'job_family': list(JobFamily.objects.filter(tenant=tenant).values_list('name', flat=True)),
    })
    _data_sheet(workbook, STAFF_SHEET, STAFF_COLUMNS, lists)
    workbook.move_sheet(STAFF_SHEET, offset=-1)
    _instructions(instructions, f'Staff import for {tenant.name}', [
        'Fill in one row per officer on the Staff sheet, then upload this file on the Import page in LSMS.',
        'Columns marked * are required. A blank cell leaves what LSMS already holds unchanged.',
        f'Officers are matched on their staff number ({terms["staff_number"]}), so you can upload an updated file at any time.',
        f'Grade here means {terms["grade"]}; Job family means {terms["job_family"]}.',
        'You can also upload an export from your HR system as it is: common header names (PF No., Job Group, Cadre, Designation...) are recognised,',
        'and any other columns are kept on the officer\'s record.',
        'LSMS checks the file first and shows what each row will do before anything is saved. Nobody is ever deleted.',
    ], [(STAFF_SHEET, STAFF_COLUMNS)])
    return _save(workbook)


def competency_template(tenant):
    workbook = Workbook()
    instructions = workbook.active
    instructions.title = 'Instructions'
    roles = JobRole.objects.filter(tenant=tenant).order_by('title').values_list('title', flat=True).distinct()
    lists = _lists_sheet(workbook, {
        'type': list(CATEGORIES),
        'job_role': list(roles),
        'required_level': [str(level) for level, _ in PROFICIENCY_LEVELS],
    })
    _data_sheet(workbook, COMPETENCY_SHEET, COMPETENCY_COLUMNS, lists)
    _data_sheet(workbook, REQUIREMENT_SHEET, REQUIREMENT_COLUMNS, lists)
    workbook.move_sheet(COMPETENCY_SHEET, offset=-1)
    workbook.move_sheet(REQUIREMENT_SHEET, offset=-1)
    _instructions(instructions, f'Competency framework import for {tenant.name}', [
        'Use this template to load a competency framework (for example a public service commission\'s) into LSMS.',
        'Competencies sheet: one row per competency. Competencies become skills that staff are assessed on.',
        'Role requirements sheet (optional): the proficiency level each post requires in each competency, which is what gaps are measured against.',
        f'Proficiency levels: {LEVEL_NAMES}.',
        'Difficulty, importance and frequency (DIF) set how much a gap in the competency matters when ranking training needs.',
        'Columns marked * are required. Competencies are matched on their code, so you can upload an updated framework later.',
        'LSMS checks the file first and shows what each row will do before anything is saved.',
    ], [(COMPETENCY_SHEET, COMPETENCY_COLUMNS), (REQUIREMENT_SHEET, REQUIREMENT_COLUMNS)])
    return _save(workbook)


TEMPLATES = {'staff': staff_template, 'competencies': competency_template}
